"""
RK Traders - Grocery Inventory & Billing System
Run:  python app.py   then open http://localhost:5000
"""
import os
import io
import json
import sqlite3
from datetime import date, datetime, timedelta
from functools import wraps

from flask import (Flask, g, render_template, request, redirect, url_for,
                   session, flash, jsonify, send_file, abort, send_from_directory)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "rktraders.db")

app = Flask(__name__)
app.config["SECRET_KEY"] = "change-this-secret-key-for-rk-traders"
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024          # 5 MB upload limit
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads", "leave")
ALLOWED_DOCS = {"pdf", "jpg", "jpeg", "png"}
LEAVE_TYPES = ["Casual", "Sick", "Earned", "Unpaid", "Other"]


@app.errorhandler(413)
def too_large(e):
    flash("File is too large. Maximum size is 5 MB.", "error")
    return redirect(request.referrer or url_for("dashboard"))

UNITS = ["pcs", "kg", "g", "ltr", "ml", "packet", "box", "dozen", "bag"]
GST_RATES = [0, 5, 12, 18, 28]
PAYMENT_MODES = ["Cash", "UPI", "Card", "Credit"]


# --------------------------------------------------------------------------
# Database helpers
# --------------------------------------------------------------------------
def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def q(sql, params=(), one=False):
    cur = get_db().execute(sql, params)
    rows = cur.fetchall()
    return (rows[0] if rows else None) if one else rows


def run(sql, params=()):
    cur = get_db().execute(sql, params)
    return cur.lastrowid


def migrate(db):
    """Add columns introduced after the first release to an existing database."""
    wanted = {
        "sales": [("prev_balance", "REAL NOT NULL DEFAULT 0"), ("credit_paid", "REAL NOT NULL DEFAULT 0")],
        "payments": [("receipt_no", "TEXT"), ("prev_balance", "REAL NOT NULL DEFAULT 0"), ("new_balance", "REAL NOT NULL DEFAULT 0")],
        "branches": [("next_receipt_seq", "INTEGER NOT NULL DEFAULT 1"), ("signature_file", "TEXT"), ("signatory_name", "TEXT")],
        "sales_extra": [],
    }
    wanted["sales"].append(("login_session_id", "INTEGER"))
    wanted["sales"].append(("offer_discount", "REAL NOT NULL DEFAULT 0"))
    wanted["sale_items"] = [("offer_name", "TEXT"), ("offer_discount", "REAL NOT NULL DEFAULT 0")]
    wanted.pop("sales_extra")
    wanted["purchases"] = [("bill_file", "TEXT"), ("bill_name", "TEXT")]
    wanted["staff_profiles"] = [("id_valid_till", "TEXT")]
    wanted["staff_documents"] = [("signed_file", "TEXT"), ("signed_name", "TEXT"), ("signed_at", "TEXT"), ("signed_by", "INTEGER")]
    wanted["users"] = [("joining_date", "TEXT"), ("can_manage_users", "INTEGER NOT NULL DEFAULT 0"),
                       ("can_view_analytics", "INTEGER NOT NULL DEFAULT 0"),
                       ("device_lock", "INTEGER NOT NULL DEFAULT 1"), ("bound_ip", "TEXT"), ("bound_at", "TEXT"),
                       ("can_reset_device", "INTEGER NOT NULL DEFAULT 0"), ("can_hr_docs", "INTEGER NOT NULL DEFAULT 0")]
    for table, cols in wanted.items():
        have = {r[1] for r in db.execute(f"PRAGMA table_info({table})")}
        for name, ddl in cols:
            if name not in have:
                db.execute(f"ALTER TABLE {table} ADD COLUMN {name} {ddl}")


def init_db():
    """Create tables and seed default data if the database is new."""
    new = not os.path.exists(DB_PATH)
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    with open(os.path.join(BASE_DIR, "schema.sql"), encoding="utf-8") as f:
        db.executescript(f.read())
    migrate(db)
    if db.execute("SELECT COUNT(*) FROM branches").fetchone()[0] == 0:
        db.execute("INSERT INTO branches(code,name,address,phone,gstin) VALUES (?,?,?,?,?)",
                   ("MB", "Main Branch", "Kolkata, West Bengal", "", ""))
        db.execute("INSERT INTO counters(branch_id,code,name) VALUES (1,'C1','Counter 1')")
    if db.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
        db.execute("INSERT INTO users(username,password_hash,full_name,role,branch_id) VALUES (?,?,?,?,NULL)",
                   ("admin", generate_password_hash("admin123"), "Administrator", "admin"))
    defaults = {
        "shop_name": "RK Traders",
        "shop_tagline": "Grocery & General Store",
        "shop_address": "Kolkata, West Bengal",
        "shop_phone": "",
        "shop_gstin": "",
        "state_name": "West Bengal",
        "state_code": "19",
        "invoice_footer": "Thank you for shopping with us! Goods once sold will not be taken back.",
        "allow_negative_stock": "0",
        "emp_id_prefix": "RKT",
        "designations": "Cashier, Senior Cashier, Store Helper, Delivery Staff, Supervisor, Store Manager, Branch Manager, Accountant",
        "id_card_validity_years": "2",
        "signature_file": "",
        "signatory_name": "",
        "offer_letter_template": (
            "Dear {name},\n\n"
            "We are pleased to offer you the position of {designation} at {company}, {branch}, with effect from {joining_date}.\n\n"
            "Your monthly salary will be Rs. {salary} ({salary_words} only), payable at the end of each month. Your working hours, "
            "leave entitlement and other service conditions will be as per the policies of {company}, as amended from time to time.\n\n"
            "Please sign and return a copy of this letter as a token of your acceptance. We look forward to having you on our team.\n\n"
            "Yours sincerely,"),
    }
    for k, v in defaults.items():
        db.execute("INSERT OR IGNORE INTO settings(key,value) VALUES (?,?)", (k, v))
    db.execute("UPDATE users SET device_lock=0 WHERE role='admin' AND bound_ip IS NULL AND device_lock=1 "
               "AND NOT EXISTS (SELECT 1 FROM settings WHERE key='device_lock_initialised')")
    db.execute("INSERT OR IGNORE INTO settings(key,value) VALUES ('device_lock_initialised','1')")
    prefix = (db.execute("SELECT value FROM settings WHERE key='emp_id_prefix'").fetchone() or ["RKT"])[0] or "RKT"
    for (uid,) in db.execute("SELECT u.id FROM users u LEFT JOIN staff_profiles p ON p.user_id=u.id "
                             "WHERE p.emp_code IS NULL OR p.emp_code='' ORDER BY u.id").fetchall():
        n = 1 + max([int(r[0][len(prefix) + 1:]) for r in db.execute("SELECT emp_code FROM staff_profiles WHERE emp_code LIKE ?", (prefix + "-%",))
                     if r[0][len(prefix) + 1:].isdigit()] or [0])
        db.execute("INSERT INTO staff_profiles(user_id,emp_code) VALUES (?,?) ON CONFLICT(user_id) DO UPDATE SET emp_code=excluded.emp_code",
                   (uid, f"{prefix}-{n:03d}"))
    if db.execute("SELECT COUNT(*) FROM leave_policies").fetchone()[0] == 0:
        for i, (t, d, m) in enumerate([("Casual", 12, 0), ("Sick", 10, 0), ("Earned", 15, 6), ("Unpaid", 0, 0), ("Other", 0, 0)]):
            db.execute("INSERT INTO leave_policies(leave_type,days_per_year,eligible_after_months,sort_order) VALUES (?,?,?,?)", (t, d, m, i))
    if db.execute("SELECT COUNT(*) FROM categories").fetchone()[0] == 0:
        for c in ["Grocery", "Rice & Pulses", "Oil & Ghee", "Spices", "Snacks", "Beverages",
                  "Dairy", "Personal Care", "Household", "Stationery"]:
            db.execute("INSERT INTO categories(name) VALUES (?)", (c,))
    db.commit()
    db.close()
    return new


def get_setting(key, default=""):
    row = q("SELECT value FROM settings WHERE key=?", (key,), one=True)
    return row["value"] if row else default


def all_settings():
    return {r["key"]: r["value"] for r in q("SELECT key,value FROM settings")}


# --------------------------------------------------------------------------
# Formatting helpers (Indian rupee format: 12,34,567.50)
# --------------------------------------------------------------------------
def inr(value, symbol=True):
    try:
        value = float(value or 0)
    except (TypeError, ValueError):
        value = 0.0
    neg = value < 0
    s = f"{abs(value):.2f}"
    int_part, dec = s.split(".")
    if len(int_part) > 3:
        last3 = int_part[-3:]
        rest = int_part[:-3]
        groups = []
        while len(rest) > 2:
            groups.insert(0, rest[-2:])
            rest = rest[:-2]
        if rest:
            groups.insert(0, rest)
        int_part = ",".join(groups) + "," + last3
    out = f"{int_part}.{dec}"
    if symbol:
        out = "₹" + out
    return ("-" if neg else "") + out


def qty_fmt(value):
    v = float(value or 0)
    return str(int(v)) if v == int(v) else f"{v:g}"


def number_to_words(n):
    """Convert an amount to Indian English words (for invoice)."""
    n = int(round(n))
    if n == 0:
        return "Zero"
    ones = ["", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten",
            "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen", "Seventeen",
            "Eighteen", "Nineteen"]
    tens = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]

    def two(x):
        if x < 20:
            return ones[x]
        return tens[x // 10] + ("" if x % 10 == 0 else " " + ones[x % 10])

    def three(x):
        h, r = divmod(x, 100)
        return ((ones[h] + " Hundred" + (" " if r else "")) if h else "") + (two(r) if r else "")

    parts = []
    crore, n = divmod(n, 10000000)
    lakh, n = divmod(n, 100000)
    thousand, n = divmod(n, 1000)
    if crore:
        parts.append(three(crore) + " Crore")
    if lakh:
        parts.append(two(lakh) + " Lakh")
    if thousand:
        parts.append(two(thousand) + " Thousand")
    if n:
        parts.append(three(n))
    return " ".join(parts)


app.jinja_env.filters["inr"] = inr
app.jinja_env.filters["qty"] = qty_fmt
app.jinja_env.filters["words"] = number_to_words


@app.context_processor
def inject_globals():
    ctx = {"settings": all_settings(), "today": date.today().isoformat(),
           "UNITS": UNITS, "GST_RATES": GST_RATES, "PAYMENT_MODES": PAYMENT_MODES}
    if session.get("user_id"):
        ctx["current_user"] = current_user()
        ctx["current_branch"] = q("SELECT * FROM branches WHERE id=?", (session.get("branch_id"),), one=True)
        ctx["current_counter"] = q("SELECT * FROM counters WHERE id=?", (session.get("counter_id"),), one=True)
        if ctx["current_user"] and ctx["current_user"]["role"] == "admin":
            ctx["all_branches"] = q("SELECT * FROM branches WHERE active=1 ORDER BY name")
        ctx["branch_count"] = q("SELECT COUNT(*) AS n FROM branches WHERE active=1", one=True)["n"]
        ctx["unread_messages"] = q("SELECT COUNT(*) AS n FROM messages WHERE to_branch_id=? AND is_read=0",
                                   (session.get("branch_id"),), one=True)["n"]
        ctx["unread_notices"] = q("SELECT COUNT(*) AS n FROM notice_recipients WHERE user_id=? AND read_at IS NULL",
                                  (session.get("user_id"),), one=True)["n"]
        ctx["popup_notices"] = unread_notice_rows(session.get("user_id")) if ctx["unread_notices"] else []
        ctx["pending_leaves"] = len(pending_leaves_for(ctx["current_user"])) if ctx["current_user"] and ctx["current_user"]["role"] != "cashier" else 0
        ctx["LEAVE_TYPES"] = [r["leave_type"] for r in q("SELECT leave_type FROM leave_policies WHERE active=1 ORDER BY sort_order, leave_type")]
        ctx["can_manage_users"] = ctx["current_user"] and (ctx["current_user"]["role"] == "admin" or
                                                           (ctx["current_user"]["role"] == "manager" and ctx["current_user"]["can_manage_users"]))
        ctx["can_view_analytics"] = ctx["current_user"] and (ctx["current_user"]["role"] == "admin" or ctx["current_user"]["can_view_analytics"])
        ctx["can_reset_device"] = ctx["current_user"] and (ctx["current_user"]["role"] == "admin" or
                                                          (ctx["current_user"]["role"] == "manager" and ctx["current_user"]["can_reset_device"]))
        ctx["can_hr_docs"] = ctx["current_user"] and (ctx["current_user"]["role"] == "admin" or
                                                     (ctx["current_user"]["role"] == "manager" and ctx["current_user"]["can_hr_docs"]))
    return ctx


def unread_notice_rows(uid, limit=10):
    return q("SELECT n.id, n.title, n.body, n.created_at, u.full_name AS sender FROM notice_recipients r "
             "JOIN notices n ON n.id=r.notice_id JOIN users u ON u.id=n.from_user_id "
             "WHERE r.user_id=? AND r.read_at IS NULL ORDER BY n.id DESC LIMIT ?", (uid, limit))


# --------------------------------------------------------------------------
# Auth
# --------------------------------------------------------------------------
def current_user():
    uid = session.get("user_id")
    return q("SELECT * FROM users WHERE id=? AND active=1", (uid,), one=True) if uid else None


def login_required(f):
    @wraps(f)
    def wrapper(*a, **kw):
        if not current_user():
            session.clear()
            return redirect(url_for("login", next=request.path))
        return f(*a, **kw)
    return wrapper


def role_required(*roles):
    def deco(f):
        @wraps(f)
        def wrapper(*a, **kw):
            u = current_user()
            if not u:
                return redirect(url_for("login"))
            if u["role"] not in roles:
                flash("You do not have permission for that page.", "error")
                return redirect(url_for("dashboard"))
            return f(*a, **kw)
        return wrapper
    return deco


def branch_id():
    return session.get("branch_id")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = q("SELECT * FROM users WHERE username=? AND active=1",
                 (request.form["username"].strip().lower(),), one=True)
        if user and check_password_hash(user["password_hash"], request.form["password"]):
            ip = request.remote_addr or "unknown"
            if user["device_lock"]:
                if user["bound_ip"] and user["bound_ip"] != ip:
                    g.audit_note = f"Device lock: account bound to {user['bound_ip']}, attempt from {ip}"
                    flash(f"Login blocked. This account is locked to the device it first logged in from ({user['bound_ip']}) "
                          f"and you are on a different device ({ip}). Ask the admin to reset your device.", "error")
                    return render_template("login.html")
                if not user["bound_ip"]:
                    run("UPDATE users SET bound_ip=?, bound_at=datetime('now','localtime') WHERE id=?", (ip, user["id"]))
                    get_db().commit()
            session.clear()
            session["user_id"] = user["id"]
            if user["branch_id"]:
                session["branch_id"] = user["branch_id"]
            else:
                first = q("SELECT id FROM branches WHERE active=1 ORDER BY id", one=True)
                session["branch_id"] = first["id"] if first else None
            run("UPDATE login_sessions SET logout_at=COALESCE(last_seen_at, login_at), end_reason='relogin' "
                "WHERE user_id=? AND logout_at IS NULL", (user["id"],))
            session["login_session_id"] = run(
                "INSERT INTO login_sessions(user_id,branch_id,ip,user_agent) VALUES (?,?,?,?)",
                (user["id"], session["branch_id"], request.remote_addr, (request.user_agent.string or "")[:200]))
            get_db().commit()
            return redirect(url_for("select_counter"))
        flash("Wrong username or password.", "error")
    return render_template("login.html")


@app.route("/logout")
def logout():
    g.audit_user, g.audit_branch = current_user(), session.get("branch_id")   # remembered for the activity log
    if session.get("login_session_id"):
        run("UPDATE login_sessions SET logout_at=datetime('now','localtime'), end_reason='logout' WHERE id=? AND logout_at IS NULL",
            (session["login_session_id"],))
        get_db().commit()
    session.clear()
    return redirect(url_for("login"))


@app.before_request
def touch_session():
    """Record activity on the audit session, at most once a minute."""
    sid = session.get("login_session_id")
    if not sid or request.endpoint == "static":
        return
    now = datetime.now()
    last = session.get("seen_ts", 0)
    if now.timestamp() - last > 60:
        run("UPDATE login_sessions SET last_seen_at=datetime('now','localtime') WHERE id=?", (sid,))
        get_db().commit()
        session["seen_ts"] = now.timestamp()


@app.route("/select-counter", methods=["GET", "POST"])
@login_required
def select_counter():
    """After login: admin picks branch, everyone picks a billing counter."""
    u = current_user()
    if request.method == "POST":
        counter = q("SELECT * FROM counters WHERE id=? AND active=1", (request.form.get("counter_id"),), one=True)
        if not counter:
            flash("Please choose a counter.", "error")
            return redirect(url_for("select_counter"))
        if u["role"] != "admin" and counter["branch_id"] != u["branch_id"]:
            flash("That counter belongs to another branch.", "error")
            return redirect(url_for("select_counter"))
        session["branch_id"] = counter["branch_id"]
        session["counter_id"] = counter["id"]
        if session.get("login_session_id"):
            run("UPDATE login_sessions SET branch_id=?, counter_id=? WHERE id=?", (counter["branch_id"], counter["id"], session["login_session_id"]))
            get_db().commit()
        return redirect(url_for("dashboard"))
    if u["role"] == "admin":
        counters = q("SELECT c.*, b.name AS branch_name FROM counters c JOIN branches b ON b.id=c.branch_id "
                     "WHERE c.active=1 AND b.active=1 ORDER BY b.name, c.code")
    else:
        counters = q("SELECT c.*, b.name AS branch_name FROM counters c JOIN branches b ON b.id=c.branch_id "
                     "WHERE c.active=1 AND c.branch_id=? ORDER BY c.code", (u["branch_id"],))
    return render_template("select_counter.html", counters=counters)


@app.route("/switch-branch/<int:bid>")
@role_required("admin")
def switch_branch(bid):
    session["branch_id"] = bid
    c = q("SELECT id FROM counters WHERE branch_id=? AND active=1 ORDER BY id", (bid,), one=True)
    session["counter_id"] = c["id"] if c else None
    return redirect(request.referrer or url_for("dashboard"))


# --------------------------------------------------------------------------
# Dashboard
# --------------------------------------------------------------------------
@app.route("/")
@login_required
def dashboard():
    b = branch_id()
    t = date.today().isoformat()
    today = q("SELECT COUNT(*) AS bills, COALESCE(SUM(grand_total),0) AS total, "
              "COALESCE(SUM(CASE WHEN payment_mode='Cash' THEN paid_amount END),0) AS cash, "
              "COALESCE(SUM(CASE WHEN payment_mode='UPI' THEN paid_amount END),0) AS upi, "
              "COALESCE(SUM(CASE WHEN payment_mode='Card' THEN paid_amount END),0) AS card, "
              "COALESCE(SUM(credit_amount),0) AS credit "
              "FROM sales WHERE branch_id=? AND date=?", (b, t), one=True)
    month_start = t[:8] + "01"
    month = q("SELECT COALESCE(SUM(grand_total),0) AS total, COUNT(*) AS bills FROM sales "
              "WHERE branch_id=? AND date>=?", (b, month_start), one=True)
    low = q("SELECT p.name, p.unit, p.low_stock_level, COALESCE(s.quantity,0) AS qty FROM products p "
            "LEFT JOIN stock s ON s.product_id=p.id AND s.branch_id=? "
            "WHERE p.active=1 AND COALESCE(s.quantity,0) <= p.low_stock_level ORDER BY qty LIMIT 10", (b,))
    low_count = q("SELECT COUNT(*) AS c FROM products p LEFT JOIN stock s ON s.product_id=p.id AND s.branch_id=? "
                  "WHERE p.active=1 AND COALESCE(s.quantity,0) <= p.low_stock_level", (b,), one=True)["c"]
    recent = q("SELECT s.*, c.name AS customer FROM sales s LEFT JOIN customers c ON c.id=s.customer_id "
               "WHERE s.branch_id=? ORDER BY s.id DESC LIMIT 8", (b,))
    top = q("SELECT si.name, SUM(si.quantity) AS qty, SUM(si.total) AS amount FROM sale_items si "
            "JOIN sales s ON s.id=si.sale_id WHERE s.branch_id=? AND s.date>=? "
            "GROUP BY si.product_id ORDER BY amount DESC LIMIT 5", (b, (date.today() - timedelta(days=7)).isoformat()))
    udhaar = q("SELECT COALESCE(SUM(credit_balance),0) AS t FROM customers", one=True)["t"]
    notices = q("SELECT n.*, u.full_name AS sender FROM notice_recipients r JOIN notices n ON n.id=r.notice_id "
                "JOIN users u ON u.id=n.from_user_id WHERE r.user_id=? AND r.read_at IS NULL ORDER BY n.id DESC LIMIT 5",
                (session["user_id"],))
    return render_template("dashboard.html", today_stats=today, month=month, low=low, low_count=low_count,
                           recent=recent, top=top, udhaar=udhaar, notices=notices)


# --------------------------------------------------------------------------
# Products & categories
# --------------------------------------------------------------------------
@app.route("/products")
@login_required
def products():
    s = request.args.get("q", "").strip()
    cat = request.args.get("cat", "")
    sql = ("SELECT p.*, c.name AS category, COALESCE(s.quantity,0) AS qty FROM products p "
           "LEFT JOIN categories c ON c.id=p.category_id "
           "LEFT JOIN stock s ON s.product_id=p.id AND s.branch_id=? WHERE p.active=1")
    params = [branch_id()]
    if s:
        sql += " AND p.name LIKE ?"
        params.append(f"%{s}%")
    if cat:
        sql += " AND p.category_id=?"
        params.append(cat)
    sql += " ORDER BY p.name"
    return render_template("products.html", products=q(sql, params), categories=q("SELECT * FROM categories ORDER BY name"),
                           search=s, cat=cat)


def product_form_values():
    f = request.form
    return (f["name"].strip(), f.get("category_id") or None, f.get("unit", "pcs"), f.get("hsn", "").strip(),
            float(f.get("gst_rate") or 0), float(f.get("purchase_price") or 0), float(f.get("selling_price") or 0),
            float(f["mrp"]) if f.get("mrp") else None, float(f.get("low_stock_level") or 5))


@app.route("/products/new", methods=["GET", "POST"])
@role_required("admin", "manager")
def product_new():
    if request.method == "POST":
        pid = run("INSERT INTO products(name,category_id,unit,hsn,gst_rate,purchase_price,selling_price,mrp,low_stock_level) "
                  "VALUES (?,?,?,?,?,?,?,?,?)", product_form_values())
        opening = float(request.form.get("opening_stock") or 0)
        if opening:
            run("INSERT INTO stock(product_id,branch_id,quantity) VALUES (?,?,?)", (pid, branch_id(), opening))
            run("INSERT INTO stock_adjustments(branch_id,product_id,change_qty,reason,date,created_by) VALUES (?,?,?,?,?,?)",
                (branch_id(), pid, opening, "Opening stock", date.today().isoformat(), session["user_id"]))
        get_db().commit()
        flash("Product added.", "success")
        if request.form.get("add_another"):
            return redirect(url_for("product_new"))
        return redirect(url_for("products"))
    return render_template("product_form.html", p=None, categories=q("SELECT * FROM categories ORDER BY name"))


@app.route("/products/<int:pid>/edit", methods=["GET", "POST"])
@role_required("admin", "manager")
def product_edit(pid):
    p = q("SELECT * FROM products WHERE id=?", (pid,), one=True) or abort(404)
    if request.method == "POST":
        run("UPDATE products SET name=?,category_id=?,unit=?,hsn=?,gst_rate=?,purchase_price=?,selling_price=?,mrp=?,"
            "low_stock_level=? WHERE id=?", product_form_values() + (pid,))
        get_db().commit()
        flash("Product updated.", "success")
        return redirect(url_for("products"))
    return render_template("product_form.html", p=p, categories=q("SELECT * FROM categories ORDER BY name"))


@app.route("/products/<int:pid>/delete", methods=["POST"])
@role_required("admin", "manager")
def product_delete(pid):
    run("UPDATE products SET active=0 WHERE id=?", (pid,))
    get_db().commit()
    flash("Product removed.", "success")
    return redirect(url_for("products"))


@app.route("/categories", methods=["GET", "POST"])
@role_required("admin", "manager")
def categories():
    if request.method == "POST":
        name = request.form["name"].strip()
        if name:
            try:
                run("INSERT INTO categories(name) VALUES (?)", (name,))
                get_db().commit()
            except sqlite3.IntegrityError:
                flash("Category already exists.", "error")
        return redirect(url_for("categories"))
    cats = q("SELECT c.*, (SELECT COUNT(*) FROM products p WHERE p.category_id=c.id AND p.active=1) AS n "
             "FROM categories c ORDER BY c.name")
    return render_template("categories.html", categories=cats)


@app.route("/categories/<int:cid>/delete", methods=["POST"])
@role_required("admin", "manager")
def category_delete(cid):
    run("UPDATE products SET category_id=NULL WHERE category_id=?", (cid,))
    run("DELETE FROM categories WHERE id=?", (cid,))
    get_db().commit()
    return redirect(url_for("categories"))


# --------------------------------------------------------------------------
# Stock (per branch)
# --------------------------------------------------------------------------
@app.route("/stock")
@login_required
def stock():
    s = request.args.get("q", "").strip()
    only_low = request.args.get("low") == "1"
    sql = ("SELECT p.*, c.name AS category, COALESCE(s.quantity,0) AS qty FROM products p "
           "LEFT JOIN categories c ON c.id=p.category_id LEFT JOIN stock s ON s.product_id=p.id AND s.branch_id=? "
           "WHERE p.active=1")
    params = [branch_id()]
    if s:
        sql += " AND p.name LIKE ?"
        params.append(f"%{s}%")
    if only_low:
        sql += " AND COALESCE(s.quantity,0) <= p.low_stock_level"
    sql += " ORDER BY p.name"
    rows = q(sql, params)
    value = sum(r["qty"] * r["purchase_price"] for r in rows)
    return render_template("stock.html", rows=rows, search=s, only_low=only_low, stock_value=value)


def change_stock(product_id, bid, delta):
    run("INSERT INTO stock(product_id,branch_id,quantity) VALUES (?,?,?) "
        "ON CONFLICT(product_id,branch_id) DO UPDATE SET quantity=quantity+excluded.quantity",
        (product_id, bid, delta))


@app.route("/stock/adjust", methods=["POST"])
@role_required("admin", "manager")
def stock_adjust():
    pid = int(request.form["product_id"])
    delta = float(request.form["change_qty"])
    reason = request.form.get("reason", "").strip() or "Manual adjustment"
    change_stock(pid, branch_id(), delta)
    run("INSERT INTO stock_adjustments(branch_id,product_id,change_qty,reason,date,created_by) VALUES (?,?,?,?,?,?)",
        (branch_id(), pid, delta, reason, date.today().isoformat(), session["user_id"]))
    get_db().commit()
    flash("Stock adjusted.", "success")
    return redirect(request.referrer or url_for("stock"))


@app.route("/stock/transfer", methods=["GET", "POST"])
@role_required("admin", "manager")
def stock_transfer():
    """Move stock from the current branch to another branch."""
    if request.method == "POST":
        to_b = int(request.form["to_branch"])
        pid = int(request.form["product_id"])
        qty = float(request.form["quantity"])
        if to_b == branch_id() or qty <= 0:
            flash("Choose a different branch and a positive quantity.", "error")
            return redirect(url_for("stock_transfer"))
        from_name = q("SELECT name FROM branches WHERE id=?", (branch_id(),), one=True)["name"]
        to_name = q("SELECT name FROM branches WHERE id=?", (to_b,), one=True)["name"]
        transfer_stock(pid, branch_id(), to_b, qty, f"Transfer {from_name} -> {to_name}")
        get_db().commit()
        flash("Stock transferred.", "success")
        return redirect(url_for("stock"))
    return render_template("stock_transfer.html",
                           branches=q("SELECT * FROM branches WHERE active=1 AND id<>? ORDER BY name", (branch_id(),)),
                           products=q("SELECT p.id,p.name,p.unit,COALESCE(s.quantity,0) AS qty FROM products p "
                                      "LEFT JOIN stock s ON s.product_id=p.id AND s.branch_id=? WHERE p.active=1 ORDER BY p.name",
                                      (branch_id(),)))


@app.route("/stock/history")
@role_required("admin", "manager")
def stock_history():
    rows = q("SELECT a.*, p.name, p.unit, u.full_name AS user FROM stock_adjustments a JOIN products p ON p.id=a.product_id "
             "LEFT JOIN users u ON u.id=a.created_by WHERE a.branch_id=? ORDER BY a.id DESC LIMIT 200", (branch_id(),))
    return render_template("stock_history.html", rows=rows)


# --------------------------------------------------------------------------
# Suppliers & customers
# --------------------------------------------------------------------------
@app.route("/suppliers", methods=["GET", "POST"])
@role_required("admin", "manager")
def suppliers():
    if request.method == "POST":
        f = request.form
        if f.get("id"):
            run("UPDATE suppliers SET name=?,phone=?,gstin=?,address=? WHERE id=?",
                (f["name"].strip(), f.get("phone", ""), f.get("gstin", ""), f.get("address", ""), f["id"]))
        else:
            run("INSERT INTO suppliers(name,phone,gstin,address) VALUES (?,?,?,?)",
                (f["name"].strip(), f.get("phone", ""), f.get("gstin", ""), f.get("address", "")))
        get_db().commit()
        return redirect(url_for("suppliers"))
    return render_template("suppliers.html", suppliers=q("SELECT * FROM suppliers ORDER BY name"))


@app.route("/customers", methods=["GET", "POST"])
@login_required
def customers():
    if request.method == "POST":
        f = request.form
        if f.get("id"):
            run("UPDATE customers SET name=?,phone=?,address=? WHERE id=?",
                (f["name"].strip(), f.get("phone", ""), f.get("address", ""), f["id"]))
        else:
            run("INSERT INTO customers(name,phone,address) VALUES (?,?,?)",
                (f["name"].strip(), f.get("phone", ""), f.get("address", "")))
        get_db().commit()
        return redirect(url_for("customers"))
    s = request.args.get("q", "").strip()
    rows = q("SELECT * FROM customers WHERE name LIKE ? OR phone LIKE ? ORDER BY name", (f"%{s}%", f"%{s}%"))
    total_credit = sum(r["credit_balance"] for r in rows)
    return render_template("customers.html", customers=rows, search=s, total_credit=total_credit)


@app.route("/customers/<int:cid>")
@login_required
def customer_detail(cid):
    c = q("SELECT * FROM customers WHERE id=?", (cid,), one=True) or abort(404)
    outstanding = q("SELECT *, credit_amount-credit_paid AS due FROM sales WHERE customer_id=? AND credit_amount-credit_paid>0.004 "
                    "ORDER BY id", (cid,))
    # ledger: credit bills are debits (customer owes more), receipts are credits
    ledger = q("SELECT date, created_at AS ts, 'bill' AS kind, id, invoice_no AS ref, grand_total AS bill_total, "
               "credit_amount AS debit, 0 AS credit, payment_mode AS note FROM sales WHERE customer_id=? "
               "UNION ALL "
               "SELECT date, date||' '||printf('%08d', id) AS ts, 'receipt' AS kind, id, receipt_no AS ref, 0, 0, amount, "
               "mode||' '||COALESCE(note,'') FROM payments WHERE customer_id=? ORDER BY date, ts", (cid, cid))
    rows, bal = [], 0.0
    for r in ledger:
        bal += r["debit"] - r["credit"]
        rows.append((r, bal))
    return render_template("customer_detail.html", c=c, outstanding=outstanding, ledger=rows,
                           pay_sale=request.args.get("sale_id"))


def next_receipt_no(bid):
    b = q("SELECT code, next_receipt_seq FROM branches WHERE id=?", (bid,), one=True)
    run("UPDATE branches SET next_receipt_seq=next_receipt_seq+1 WHERE id=?", (bid,))
    today = date.today()
    fy_start = today.year if today.month >= 4 else today.year - 1
    return f"{b['code']}/RCPT/{str(fy_start)[2:]}{str(fy_start + 1)[2:]}/{b['next_receipt_seq']:05d}"


@app.route("/customers/<int:cid>/payment", methods=["POST"])
@login_required
def customer_payment(cid):
    """Receive money against udhaar. Creates a numbered receipt and applies it to unpaid invoices."""
    c = q("SELECT * FROM customers WHERE id=?", (cid,), one=True) or abort(404)
    amt = round(float(request.form.get("amount") or 0), 2)
    prev = float(c["credit_balance"])
    if amt <= 0:
        flash("Enter an amount greater than zero.", "error")
        return redirect(url_for("customer_detail", cid=cid))
    if amt > prev + 0.004:
        flash(f"Amount is more than the due balance of {inr(prev)}.", "error")
        return redirect(url_for("customer_detail", cid=cid))
    mode = request.form.get("mode", "Cash")
    receipt_no = next_receipt_no(branch_id())
    pid = run("INSERT INTO payments(customer_id,branch_id,date,amount,mode,note,created_by,receipt_no,prev_balance,new_balance) "
              "VALUES (?,?,?,?,?,?,?,?,?,?)",
              (cid, branch_id(), request.form.get("date") or date.today().isoformat(), amt, mode,
               request.form.get("note", "").strip(), session["user_id"], receipt_no, prev, round(prev - amt, 2)))
    # apply to a chosen invoice first, then oldest unpaid invoices
    remaining = amt
    chosen = request.form.get("sale_id")
    order = []
    if chosen:
        order.append(int(chosen))
    order += [r["id"] for r in q("SELECT id FROM sales WHERE customer_id=? AND credit_amount-credit_paid>0.004 ORDER BY id", (cid,))
              if r["id"] not in order]
    for sid in order:
        if remaining <= 0.004:
            break
        sale = q("SELECT credit_amount-credit_paid AS due FROM sales WHERE id=? AND customer_id=?", (sid, cid), one=True)
        if not sale or sale["due"] <= 0.004:
            continue
        a = round(min(remaining, sale["due"]), 2)
        run("INSERT INTO payment_allocations(payment_id,sale_id,amount) VALUES (?,?,?)", (pid, sid, a))
        run("UPDATE sales SET credit_paid=credit_paid+? WHERE id=?", (a, sid))
        remaining = round(remaining - a, 2)
    run("UPDATE customers SET credit_balance=credit_balance-? WHERE id=?", (amt, cid))
    get_db().commit()
    flash(f"Receipt {receipt_no} saved for {inr(amt)}.", "success")
    return redirect(url_for("receipt_detail", rid=pid, print=1))


@app.route("/receipts")
@login_required
def receipts():
    f = request.args.get("from") or date.today().replace(day=1).isoformat()
    t = request.args.get("to") or date.today().isoformat()
    rows = q("SELECT p.*, c.name AS customer, u.full_name AS user FROM payments p JOIN customers c ON c.id=p.customer_id "
             "LEFT JOIN users u ON u.id=p.created_by WHERE p.branch_id=? AND p.date BETWEEN ? AND ? ORDER BY p.id DESC",
             (branch_id(), f, t))
    return render_template("receipts.html", receipts=rows, date_from=f, date_to=t, total=sum(r["amount"] for r in rows))


@app.route("/receipts/<int:rid>")
@login_required
def receipt_detail(rid):
    p = q("SELECT p.*, c.name AS customer, c.phone AS customer_phone, c.address AS customer_address, "
          "b.name AS branch_name, b.address AS branch_address, b.phone AS branch_phone, b.gstin AS branch_gstin, "
          "u.full_name AS received_by FROM payments p JOIN customers c ON c.id=p.customer_id "
          "LEFT JOIN branches b ON b.id=p.branch_id LEFT JOIN users u ON u.id=p.created_by WHERE p.id=?", (rid,), one=True) or abort(404)
    allocs = q("SELECT a.amount, s.id AS sale_id, s.invoice_no, s.date, s.grand_total, s.credit_amount, s.credit_paid, "
               "s.credit_amount-s.credit_paid AS still_due FROM payment_allocations a JOIN sales s ON s.id=a.sale_id "
               "WHERE a.payment_id=? ORDER BY s.id", (rid,))
    sig_path, sig_name, sig_src = signature_for(p["branch_id"])
    return render_template("receipt.html", p=p, allocs=allocs, autoprint=request.args.get("print") == "1",
                           sig_url=(url_for("branch_signature", bid=p["branch_id"] or 0) if sig_path else None), sig_name=sig_name)


@app.route("/receipts/<int:rid>/delete", methods=["POST"])
@role_required("admin")
def receipt_delete(rid):
    """Cancel a receipt: invoices become due again and the customer balance goes back up."""
    p = q("SELECT * FROM payments WHERE id=?", (rid,), one=True) or abort(404)
    for a in q("SELECT * FROM payment_allocations WHERE payment_id=?", (rid,)):
        run("UPDATE sales SET credit_paid=credit_paid-? WHERE id=?", (a["amount"], a["sale_id"]))
    run("DELETE FROM payment_allocations WHERE payment_id=?", (rid,))
    run("DELETE FROM payments WHERE id=?", (rid,))
    run("UPDATE customers SET credit_balance=credit_balance+? WHERE id=?", (p["amount"], p["customer_id"]))
    get_db().commit()
    flash(f"Receipt {p['receipt_no']} cancelled.", "success")
    return redirect(url_for("customer_detail", cid=p["customer_id"]))


# --------------------------------------------------------------------------
# Purchases (stock in)
# --------------------------------------------------------------------------
@app.route("/purchases")
@role_required("admin", "manager")
def purchases():
    rows = q("SELECT p.*, s.name AS supplier, u.full_name AS user FROM purchases p "
             "LEFT JOIN suppliers s ON s.id=p.supplier_id LEFT JOIN users u ON u.id=p.created_by "
             "WHERE p.branch_id=? ORDER BY p.id DESC LIMIT 200", (branch_id(),))
    return render_template("purchases.html", purchases=rows)


PURCHASE_BILL_DIR = os.path.join(BASE_DIR, "uploads", "purchases")


def save_purchase_bill(file):
    """Store an uploaded supplier bill (PDF/JPG/PNG). Returns (stored_name, original_name) or (None, None)."""
    if not file or not file.filename:
        return None, None
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_DOCS:
        raise ValueError("Supplier bill must be a PDF, JPG or PNG file.")
    os.makedirs(PURCHASE_BILL_DIR, exist_ok=True)
    stored = f"bill_{int(datetime.now().timestamp())}_{secure_filename(file.filename)[:60]}"
    file.save(os.path.join(PURCHASE_BILL_DIR, stored))
    return stored, file.filename


@app.route("/purchases/new", methods=["GET", "POST"])
@role_required("admin", "manager")
def purchase_new():
    if request.method == "POST":
        if request.is_json:
            data = request.get_json(force=True)
            bill_file = request.files.get("bill")
        else:                                            # multipart from the form (items as JSON text + optional bill file)
            data = dict(request.form)
            try:
                data["items"] = json.loads(request.form.get("items") or "[]")
            except ValueError:
                return jsonify(ok=False, error="Bad item data."), 400
            bill_file = request.files.get("bill")
        items = [i for i in data.get("items", []) if i.get("product_id") and float(i.get("quantity") or 0) > 0]
        if not items:
            return jsonify(ok=False, error="Add at least one item."), 400
        try:
            stored, orig = save_purchase_bill(bill_file)
        except ValueError as e:
            return jsonify(ok=False, error=str(e)), 400
        if not stored:
            return jsonify(ok=False, error="Please attach the supplier's bill (PDF or photo). A purchase cannot be saved without it."), 400
        total = sum(float(i["quantity"]) * float(i["unit_cost"]) for i in items)
        pid = run("INSERT INTO purchases(branch_id,supplier_id,invoice_no,date,total,note,created_by,bill_file,bill_name) VALUES (?,?,?,?,?,?,?,?,?)",
                  (branch_id(), data.get("supplier_id") or None, data.get("invoice_no", ""),
                   data.get("date") or date.today().isoformat(), total, data.get("note", ""), session["user_id"], stored, orig))
        for i in items:
            qty, cost = float(i["quantity"]), float(i["unit_cost"])
            run("INSERT INTO purchase_items(purchase_id,product_id,quantity,unit_cost,total) VALUES (?,?,?,?,?)",
                (pid, i["product_id"], qty, cost, qty * cost))
            change_stock(int(i["product_id"]), branch_id(), qty)
            if i.get("update_cost"):
                run("UPDATE products SET purchase_price=? WHERE id=?", (cost, i["product_id"]))
            if i.get("selling_price"):
                run("UPDATE products SET selling_price=? WHERE id=?", (float(i["selling_price"]), i["product_id"]))
        get_db().commit()
        return jsonify(ok=True, id=pid)
    return render_template("purchase_form.html", suppliers=q("SELECT * FROM suppliers ORDER BY name"))


@app.route("/purchases/<int:pid>/bill", methods=["GET", "POST"])
@role_required("admin", "manager")
def purchase_bill(pid):
    p = q("SELECT * FROM purchases WHERE id=?", (pid,), one=True) or abort(404)
    if request.method == "POST":
        try:
            stored, orig = save_purchase_bill(request.files.get("bill"))
        except ValueError as e:
            flash(str(e), "error")
            return redirect(url_for("purchase_detail", pid=pid))
        if not stored:
            flash("Choose a file first.", "error")
            return redirect(url_for("purchase_detail", pid=pid))
        if p["bill_file"]:
            try:
                os.remove(os.path.join(PURCHASE_BILL_DIR, p["bill_file"]))
            except OSError:
                pass
        run("UPDATE purchases SET bill_file=?, bill_name=? WHERE id=?", (stored, orig, pid))
        get_db().commit()
        flash("Supplier bill attached.", "success")
        return redirect(url_for("purchase_detail", pid=pid))
    if not p["bill_file"]:
        abort(404)
    return send_from_directory(PURCHASE_BILL_DIR, p["bill_file"], as_attachment=False, download_name=p["bill_name"])


@app.route("/purchases/<int:pid>")
@role_required("admin", "manager")
def purchase_detail(pid):
    p = q("SELECT p.*, s.name AS supplier, s.phone AS supplier_phone FROM purchases p LEFT JOIN suppliers s ON s.id=p.supplier_id "
          "WHERE p.id=?", (pid,), one=True) or abort(404)
    items = q("SELECT pi.*, pr.name, pr.unit FROM purchase_items pi JOIN products pr ON pr.id=pi.product_id WHERE purchase_id=?", (pid,))
    return render_template("purchase_detail.html", p=p, items=items)


# --------------------------------------------------------------------------
# Offers
# --------------------------------------------------------------------------
OFFER_TYPES = {"product_pct": "% off a product", "product_price": "Special price for a product", "buy_x_get_y": "Buy X get Y free",
               "category_pct": "% off a category", "bill_pct": "% off the bill above an amount", "bill_flat": "Flat amount off the bill above an amount"}


def active_offers(bid, on=None):
    on = on or date.today().isoformat()
    return [dict(r) for r in q("SELECT o.*, p.name AS product_name, c.name AS category_name FROM offers o "
                               "LEFT JOIN products p ON p.id=o.product_id LEFT JOIN categories c ON c.id=o.category_id "
                               "WHERE o.active=1 AND o.from_date<=? AND o.to_date>=? AND (o.branch_id IS NULL OR o.branch_id=?) ORDER BY o.id",
                               (on, on, bid))]


def line_offer(line, offers):
    """Best offer for one cart line. line: dict(product_id, category_id, qty, price). Returns (name, discount)."""
    best = ("", 0.0)
    gross = line["qty"] * line["price"]
    for o in offers:
        d = 0.0
        if o["type"] == "product_pct" and o["product_id"] == line["product_id"]:
            d = gross * o["value"] / 100
        elif o["type"] == "product_price" and o["product_id"] == line["product_id"] and o["value"] < line["price"]:
            d = (line["price"] - o["value"]) * line["qty"]
        elif o["type"] == "buy_x_get_y" and o["product_id"] == line["product_id"] and o["buy_qty"] and o["free_qty"]:
            groups = int(line["qty"] // (o["buy_qty"] + o["free_qty"]))
            d = groups * o["free_qty"] * line["price"]
        elif o["type"] == "category_pct" and o["category_id"] and o["category_id"] == line["category_id"]:
            d = gross * o["value"] / 100
        d = round(min(d, gross), 2)
        if d > best[1]:
            best = (o["name"], d)
    return best


def bill_offer(subtotal_after_lines, offers):
    best = ("", 0.0)
    for o in offers:
        if o["type"] not in ("bill_pct", "bill_flat") or subtotal_after_lines < (o["min_bill"] or 0) or subtotal_after_lines <= 0:
            continue
        d = subtotal_after_lines * o["value"] / 100 if o["type"] == "bill_pct" else o["value"]
        d = round(min(d, subtotal_after_lines), 2)
        if d > best[1]:
            best = (o["name"], d)
    return best


def offer_form_values(f, u):
    t = f.get("type")
    if t not in OFFER_TYPES:
        raise ValueError("Choose an offer type.")
    name = f.get("name", "").strip()
    if not name:
        raise ValueError("Give the offer a name.")
    d1, d2 = f.get("from_date"), f.get("to_date")
    if not d1 or not d2 or d2 < d1:
        raise ValueError("Enter a valid date range.")
    value = float(f.get("value") or 0)
    product_id = int(f["product_id"]) if f.get("product_id") else None
    category_id = int(f["category_id"]) if f.get("category_id") else None
    buy_qty = float(f.get("buy_qty") or 0) or None
    free_qty = float(f.get("free_qty") or 0) or None
    min_bill = float(f.get("min_bill") or 0)
    if t in ("product_pct", "product_price", "buy_x_get_y") and not product_id:
        raise ValueError("Pick a product.")
    if t == "category_pct" and not category_id:
        raise ValueError("Pick a category.")
    if t in ("product_pct", "category_pct", "bill_pct") and not (0 < value <= 100):
        raise ValueError("Percentage must be between 1 and 100.")
    if t in ("product_price", "bill_flat") and value <= 0:
        raise ValueError("Enter an amount greater than zero.")
    if t == "buy_x_get_y" and not (buy_qty and free_qty):
        raise ValueError("Enter buy quantity and free quantity.")
    if u["role"] == "admin":
        branch = int(f["branch_id"]) if f.get("branch_id") else None
    else:
        branch = u["branch_id"]
    return (name, t, product_id, category_id, value, buy_qty, free_qty, min_bill, d1, d2, branch, 1 if f.get("active", "1") else 0)


@app.route("/offers")
@role_required("admin", "manager")
def offers():
    u = current_user()
    today = date.today().isoformat()
    sql = ("SELECT o.*, p.name AS product_name, c.name AS category_name, b.name AS branch_name, u.full_name AS creator FROM offers o "
           "LEFT JOIN products p ON p.id=o.product_id LEFT JOIN categories c ON c.id=o.category_id LEFT JOIN branches b ON b.id=o.branch_id "
           "LEFT JOIN users u ON u.id=o.created_by ")
    rows = q(sql + "ORDER BY o.active DESC, o.to_date DESC, o.id DESC") if u["role"] == "admin" else \
        q(sql + "WHERE o.branch_id IS NULL OR o.branch_id=? ORDER BY o.active DESC, o.to_date DESC, o.id DESC", (u["branch_id"],))
    return render_template("offers.html", offers=rows, today=today, OFFER_TYPES=OFFER_TYPES)


@app.route("/offers/new", methods=["GET", "POST"])
@app.route("/offers/<int:oid>/edit", methods=["GET", "POST"])
@role_required("admin", "manager")
def offer_form(oid=None):
    u = current_user()
    o = None
    if oid:
        o = q("SELECT o.*, p.name AS product_name FROM offers o LEFT JOIN products p ON p.id=o.product_id WHERE o.id=?", (oid,), one=True) or abort(404)
        if u["role"] != "admin" and o["branch_id"] != u["branch_id"]:
            flash("You can only edit offers of your own branch.", "error")
            return redirect(url_for("offers"))
    if request.method == "POST":
        try:
            vals = offer_form_values(request.form, u)
        except ValueError as e:
            flash(str(e), "error")
            return redirect(request.url)
        if oid:
            run("UPDATE offers SET name=?,type=?,product_id=?,category_id=?,value=?,buy_qty=?,free_qty=?,min_bill=?,from_date=?,to_date=?,"
                "branch_id=?,active=? WHERE id=?", vals + (oid,))
            flash("Offer updated.", "success")
        else:
            run("INSERT INTO offers(name,type,product_id,category_id,value,buy_qty,free_qty,min_bill,from_date,to_date,branch_id,active,created_by) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)", vals + (u["id"],))
            flash("Offer created. It applies automatically at billing.", "success")
        get_db().commit()
        return redirect(url_for("offers"))
    return render_template("offer_form.html", o=o, OFFER_TYPES=OFFER_TYPES, categories=q("SELECT * FROM categories ORDER BY name"),
                           branches_list=q("SELECT * FROM branches WHERE active=1 ORDER BY name"))


@app.route("/offers/<int:oid>/toggle", methods=["POST"])
@role_required("admin", "manager")
def offer_toggle(oid):
    u = current_user()
    o = q("SELECT * FROM offers WHERE id=?", (oid,), one=True) or abort(404)
    if u["role"] != "admin" and o["branch_id"] != u["branch_id"]:
        abort(403)
    run("UPDATE offers SET active=1-active WHERE id=?", (oid,))
    get_db().commit()
    return redirect(url_for("offers"))


@app.route("/offers/<int:oid>/delete", methods=["POST"])
@role_required("admin", "manager")
def offer_delete(oid):
    u = current_user()
    o = q("SELECT * FROM offers WHERE id=?", (oid,), one=True) or abort(404)
    if u["role"] != "admin" and o["branch_id"] != u["branch_id"]:
        abort(403)
    run("DELETE FROM offers WHERE id=?", (oid,))
    get_db().commit()
    flash("Offer deleted.", "success")
    return redirect(url_for("offers"))


@app.route("/api/offers")
@login_required
def api_offers():
    return jsonify(active_offers(branch_id()))


# --------------------------------------------------------------------------
# Billing (point of sale)
# --------------------------------------------------------------------------
@app.route("/billing")
@login_required
def billing():
    if not session.get("counter_id"):
        flash("Please select a billing counter first.", "error")
        return redirect(url_for("select_counter"))
    return render_template("billing.html")


@app.route("/api/products/search")
@login_required
def api_product_search():
    s = request.args.get("q", "").strip()
    rows = q("SELECT p.id,p.name,p.unit,p.selling_price,p.mrp,p.gst_rate,p.hsn,p.category_id,COALESCE(s.quantity,0) AS qty "
             "FROM products p LEFT JOIN stock s ON s.product_id=p.id AND s.branch_id=? "
             "WHERE p.active=1 AND p.name LIKE ? ORDER BY CASE WHEN p.name LIKE ? THEN 0 ELSE 1 END, p.name LIMIT 15",
             (branch_id(), f"%{s}%", f"{s}%"))
    return jsonify([dict(r) for r in rows])


@app.route("/api/customers/search")
@login_required
def api_customer_search():
    s = request.args.get("q", "").strip()
    rows = q("SELECT id,name,phone,credit_balance FROM customers WHERE name LIKE ? OR phone LIKE ? ORDER BY name LIMIT 10",
             (f"%{s}%", f"%{s}%"))
    return jsonify([dict(r) for r in rows])


@app.route("/api/customers", methods=["POST"])
@login_required
def api_customer_create():
    d = request.get_json(force=True)
    if not d.get("name", "").strip():
        return jsonify(ok=False, error="Name required"), 400
    cid = run("INSERT INTO customers(name,phone,address) VALUES (?,?,?)", (d["name"].strip(), d.get("phone", ""), ""))
    get_db().commit()
    return jsonify(ok=True, id=cid, name=d["name"].strip(), phone=d.get("phone", ""), credit_balance=0)


def next_invoice_no(bid, cid):
    b = q("SELECT code FROM branches WHERE id=?", (bid,), one=True)
    c = q("SELECT code, next_seq FROM counters WHERE id=?", (cid,), one=True)
    seq = c["next_seq"]
    run("UPDATE counters SET next_seq=next_seq+1 WHERE id=?", (cid,))
    today = date.today()
    fy_start = today.year if today.month >= 4 else today.year - 1          # Indian financial year April-March
    fy_label = f"{str(fy_start)[2:]}{str(fy_start + 1)[2:]}"
    return f"{b['code']}/{c['code']}/{fy_label}/{seq:05d}"


@app.route("/api/sales", methods=["POST"])
@login_required
def api_sale_create():
    d = request.get_json(force=True)
    bid, cid = branch_id(), session.get("counter_id")
    if not cid:
        return jsonify(ok=False, error="No billing counter selected."), 400
    items = d.get("items", [])
    if not items:
        return jsonify(ok=False, error="Bill is empty."), 400
    allow_neg = get_setting("allow_negative_stock") == "1"
    discount = float(d.get("discount") or 0)
    mode = d.get("payment_mode", "Cash")
    customer_id = d.get("customer_id") or None

    # Validate stock and compute totals (list price, then offers, then manual discount)
    offers_now = active_offers(bid)
    subtotal = 0.0
    lines = []
    for it in items:
        p = q("SELECT p.*, COALESCE(s.quantity,0) AS qty FROM products p LEFT JOIN stock s ON s.product_id=p.id AND s.branch_id=? "
              "WHERE p.id=?", (bid, it["product_id"]), one=True)
        if not p:
            return jsonify(ok=False, error="Unknown product."), 400
        qty = float(it["quantity"])
        price = float(it.get("unit_price") or p["selling_price"])
        if qty <= 0:
            return jsonify(ok=False, error=f"Invalid quantity for {p['name']}."), 400
        if not allow_neg and p["qty"] < qty:
            return jsonify(ok=False, error=f"Only {qty_fmt(p['qty'])} {p['unit']} of {p['name']} in stock."), 400
        line_total = round(qty * price, 2)
        subtotal += line_total
        oname, odisc = line_offer({"product_id": p["id"], "category_id": p["category_id"], "qty": qty, "price": price}, offers_now)
        lines.append((p, qty, price, line_total, oname, odisc))

    line_offer_total = round(sum(l[5] for l in lines), 2)
    after_lines = round(subtotal - line_offer_total, 2)
    bname, bdisc = bill_offer(after_lines, offers_now)
    offer_total = round(line_offer_total + bdisc, 2)
    if discount < 0 or discount > after_lines - bdisc:
        return jsonify(ok=False, error="Invalid discount."), 400
    # bill offer + manual discount are spread across lines in proportion to their value after line offers
    spread = bdisc + discount
    disc_ratio = (after_lines - spread) / after_lines if after_lines else 1

    taxable_total = tax_total = 0.0
    item_rows = []
    for p, qty, price, line_total, oname, odisc in lines:
        net = round((line_total - odisc) * disc_ratio, 2)      # after all discounts, GST inclusive
        rate = float(p["gst_rate"])
        taxable = round(net / (1 + rate / 100), 2)
        tax = round(net - taxable, 2)
        taxable_total += taxable
        tax_total += tax
        if bdisc and not oname:
            oname = bname
        item_rows.append((p, qty, price, rate, taxable, tax, net, oname, odisc))

    exact = subtotal - offer_total - discount
    grand = float(int(exact + 0.5))                           # round to nearest rupee, .50 rounds up
    round_off = round(grand - exact, 2)
    cgst = round(tax_total / 2, 2)
    sgst = round(tax_total - cgst, 2)

    paid = float(d.get("paid_amount")) if d.get("paid_amount") not in (None, "") else grand
    credit = 0.0
    prev_balance = 0.0
    if mode == "Credit":
        if not customer_id:
            return jsonify(ok=False, error="Select a customer for a credit bill."), 400
        cust = q("SELECT credit_balance FROM customers WHERE id=?", (customer_id,), one=True)
        if not cust:
            return jsonify(ok=False, error="Customer not found."), 400
        prev_balance = float(cust["credit_balance"])
        credit = round(grand - paid, 2)
        if credit < 0:
            credit = 0.0
    else:
        paid = grand

    inv = next_invoice_no(bid, cid)
    sid = run("INSERT INTO sales(invoice_no,branch_id,counter_id,customer_id,date,subtotal,discount,taxable,cgst,sgst,"
              "round_off,grand_total,payment_mode,paid_amount,credit_amount,prev_balance,created_by,login_session_id,offer_discount) "
              "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
              (inv, bid, cid, customer_id, date.today().isoformat(), subtotal, discount, round(taxable_total, 2),
               cgst, sgst, round_off, grand, mode, paid, credit, prev_balance, session["user_id"], session.get("login_session_id"), offer_total))
    for p, qty, price, rate, taxable, tax, net, oname, odisc in item_rows:
        run("INSERT INTO sale_items(sale_id,product_id,name,hsn,unit,quantity,unit_price,cost_price,gst_rate,taxable,tax,total,offer_name,offer_discount) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (sid, p["id"], p["name"], p["hsn"], p["unit"], qty, price, p["purchase_price"], rate, taxable, tax, net, oname or None, odisc))
        change_stock(p["id"], bid, -qty)
    if credit:
        run("UPDATE customers SET credit_balance=credit_balance+? WHERE id=?", (credit, customer_id))
    get_db().commit()
    return jsonify(ok=True, id=sid, invoice_no=inv, grand_total=grand, offer_discount=offer_total)


@app.route("/sales")
@login_required
def sales():
    f = request.args.get("from") or date.today().isoformat()
    t = request.args.get("to") or date.today().isoformat()
    rows = q("SELECT s.*, c.name AS customer, u.full_name AS user, co.code AS counter FROM sales s "
             "LEFT JOIN customers c ON c.id=s.customer_id LEFT JOIN users u ON u.id=s.created_by "
             "LEFT JOIN counters co ON co.id=s.counter_id "
             "WHERE s.branch_id=? AND s.date BETWEEN ? AND ? ORDER BY s.id DESC", (branch_id(), f, t))
    total = sum(r["grand_total"] for r in rows)
    return render_template("sales.html", sales=rows, date_from=f, date_to=t, total=total)


@app.route("/sales/<int:sid>")
@login_required
def sale_detail(sid):
    s = q("SELECT s.*, c.name AS customer, c.phone AS customer_phone, c.address AS customer_address, "
          "b.name AS branch_name, b.address AS branch_address, b.phone AS branch_phone, b.gstin AS branch_gstin, "
          "co.name AS counter_name, u.full_name AS cashier FROM sales s "
          "LEFT JOIN customers c ON c.id=s.customer_id JOIN branches b ON b.id=s.branch_id "
          "JOIN counters co ON co.id=s.counter_id LEFT JOIN users u ON u.id=s.created_by WHERE s.id=?", (sid,), one=True) or abort(404)
    items = q("SELECT * FROM sale_items WHERE sale_id=? ORDER BY id", (sid,))
    tax_summary = q("SELECT gst_rate, SUM(taxable) AS taxable, SUM(tax) AS tax FROM sale_items WHERE sale_id=? "
                    "GROUP BY gst_rate ORDER BY gst_rate", (sid,))
    receipts_applied = q("SELECT a.amount, p.id AS payment_id, p.receipt_no, p.date, p.mode FROM payment_allocations a "
                         "JOIN payments p ON p.id=a.payment_id WHERE a.sale_id=? ORDER BY p.id", (sid,))
    sig_path, sig_name, sig_src = signature_for(s["branch_id"])
    return render_template("invoice.html", s=s, items=items, tax_summary=tax_summary, receipts_applied=receipts_applied,
                           autoprint=request.args.get("print") == "1",
                           sig_url=(url_for("branch_signature", bid=s["branch_id"]) if sig_path else None), sig_name=sig_name)


@app.route("/sales/<int:sid>/delete", methods=["POST"])
@role_required("admin")
def sale_delete(sid):
    """Cancel a bill: stock goes back, customer credit reversed."""
    s = q("SELECT * FROM sales WHERE id=?", (sid,), one=True) or abort(404)
    if s["credit_paid"]:
        flash(f"Bill {s['invoice_no']} cannot be cancelled: a payment receipt is already applied to it. Cancel that receipt first.", "error")
        return redirect(url_for("sales"))
    for it in q("SELECT * FROM sale_items WHERE sale_id=?", (sid,)):
        change_stock(it["product_id"], s["branch_id"], it["quantity"])
    if s["credit_amount"] and s["customer_id"]:
        run("UPDATE customers SET credit_balance=credit_balance-? WHERE id=?", (s["credit_amount"], s["customer_id"]))
    run("DELETE FROM sale_items WHERE sale_id=?", (sid,))
    run("DELETE FROM sales WHERE id=?", (sid,))
    get_db().commit()
    flash(f"Bill {s['invoice_no']} cancelled and stock restored.", "success")
    return redirect(url_for("sales"))


# --------------------------------------------------------------------------
# Export engine: Excel / CSV / PDF with company + branch header
# --------------------------------------------------------------------------
def _fmt_cell(v, for_excel=False):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, float):
        return round(v, 2) if for_excel else inr(v, symbol=False)
    return v


def export_header(title, period):
    st = all_settings()
    u = current_user()
    b = q("SELECT * FROM branches WHERE id=?", (session.get("branch_id"),), one=True)
    company = [st.get("shop_name", ""), st.get("shop_tagline", ""), st.get("shop_address", "")]
    if st.get("shop_phone"):
        company.append(f"Phone: {st['shop_phone']}")
    if st.get("shop_gstin"):
        company.append(f"GSTIN: {st['shop_gstin']}  State: {st.get('state_name', '')} ({st.get('state_code', '')})")
    branch = []
    if b:
        branch = [f"Branch: {b['name']} ({b['code']})"]
        if b["address"]:
            branch.append(b["address"])
        if b["phone"]:
            branch.append(f"Phone: {b['phone']}")
        if b["gstin"]:
            branch.append(f"Branch GSTIN: {b['gstin']}")
    meta = [title, f"Period: {period}" if period else "",
            f"Generated by {u['full_name'] if u else '-'} on {datetime.now().strftime('%d-%m-%Y %I:%M %p')}"]
    return [x for x in company if x], branch, [x for x in meta if x]


def export_response(title, period, rows, fmt, columns=None, filename="report"):
    """rows: list of sqlite Rows/dicts. columns: list of (key, label) or None to use all keys."""
    rows = [dict(r) for r in rows]
    if columns is None:
        keys = list(rows[0].keys()) if rows else []
        columns = [(k, k.replace("_", " ").title()) for k in keys]
    company, branch, meta = export_header(title, period)
    if not columns:
        columns = [("_", "Result")]
    labels = [c[1] for c in columns]
    fname = f"{filename}_{date.today().isoformat()}"
    buf = io.BytesIO()

    if fmt == "csv":
        import csv
        out = io.StringIO()
        w = csv.writer(out)
        for line in company + branch + meta:
            w.writerow([line])
        w.writerow([])
        w.writerow(labels)
        for r in rows:
            w.writerow([_fmt_cell(r.get(k)) for k, _ in columns])
        buf.write(("\ufeff" + out.getvalue()).encode("utf-8"))
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"{fname}.csv", mimetype="text/csv")

    if fmt == "pdf":
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        font = "Helvetica"
        try:
            pdfmetrics.registerFont(TTFont("Arial", r"C:\Windows\Fonts\arial.ttf"))
            pdfmetrics.registerFont(TTFont("Arial-Bold", r"C:\Windows\Fonts\arialbd.ttf"))
            font = "Arial"
        except Exception:
            pass
        bold = font + "-Bold" if font == "Arial" else "Helvetica-Bold"
        pagesize = landscape(A4) if len(columns) > 6 else A4
        doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                                title=title)
        ss = getSampleStyleSheet()
        h1 = ParagraphStyle("h1", parent=ss["Title"], fontName=bold, fontSize=16, alignment=0, spaceAfter=2)
        small = ParagraphStyle("small", parent=ss["Normal"], fontName=font, fontSize=9, leading=11)
        bsty = ParagraphStyle("b", parent=small, fontName=bold)
        head = [Paragraph(company[0], h1)] + [Paragraph(line, small) for line in company[1:]]
        logo = logo_drawing(40)
        if logo:
            lt = Table([[logo, head]], colWidths=[48, None])
            lt.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0)]))
            story = [lt]
        else:
            story = head
        for line in branch:
            story.append(Paragraph(line, bsty if line.startswith("Branch:") else small))
        story.append(Spacer(1, 4))
        story.append(Paragraph(meta[0], ParagraphStyle("t", parent=ss["Heading2"], fontName=bold, fontSize=13, spaceBefore=4, spaceAfter=2)))
        for line in meta[1:]:
            story.append(Paragraph(line, small))
        story.append(Spacer(1, 6))
        cell = ParagraphStyle("cell", parent=small, fontSize=8, leading=10)
        data = [[Paragraph(str(l), ParagraphStyle("hd", parent=cell, fontName=bold)) for l in labels]]
        numeric_cols = set()
        for r in rows:
            line = []
            for ci, (k, _) in enumerate(columns):
                v = r.get(k)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    numeric_cols.add(ci)
                line.append(Paragraph(str(_fmt_cell(v)), cell))
            data.append(line)
        if len(data) == 1:
            data.append([Paragraph("No data", cell)] + [""] * (len(labels) - 1))
        avail = pagesize[0] - 24 * mm
        tbl = Table(data, repeatRows=1, colWidths=[avail / len(labels)] * len(labels))
        style = [("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6f4f1")),
                 ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")])]
        for ci in numeric_cols:
            style.append(("ALIGN", (ci, 1), (ci, -1), "RIGHT"))
        tbl.setStyle(TableStyle(style))
        story.append(tbl)
        story.append(Spacer(1, 8))
        story.append(Paragraph(f"{len(rows)} row(s)", small))

        def footer(canvas, d):
            canvas.saveState()
            canvas.setFont(font, 8)
            canvas.drawRightString(pagesize[0] - 12 * mm, 7 * mm, f"Page {d.page}")
            canvas.drawString(12 * mm, 7 * mm, f"{company[0]} - {title}")
            canvas.restoreState()
        doc.build(story, onFirstPage=footer, onLaterPages=footer)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"{fname}.pdf", mimetype="application/pdf")

    # default: Excel
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = title[:28]
    ws.append([company[0]])
    ws["A1"].font = Font(bold=True, size=14)
    for line in company[1:] + branch:
        ws.append([line])
        if line.startswith("Branch:"):
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([])
    ws.append([meta[0]])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    for line in meta[1:]:
        ws.append([line])
    ws.append([])
    ws.append(labels)
    hdr_row = ws.max_row
    for c in ws[hdr_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E6F4F1")
    for r in rows:
        ws.append([_fmt_cell(r.get(k), for_excel=True) for k, _ in columns])
    for ci, (k, label) in enumerate(columns, start=1):
        width = max([len(str(label))] + [len(str(r.get(k) or "")) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(ci)].width = min(max(width, 8), 60)
        for row in ws.iter_rows(min_row=hdr_row + 1, min_col=ci, max_col=ci):
            for c in row:
                if isinstance(c.value, (int, float)):
                    c.alignment = Alignment(horizontal="right")
                    if isinstance(c.value, float):
                        c.number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{fname}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------
# Reports
# --------------------------------------------------------------------------
def report_data(kind, f, t, bid):
    if kind == "daily":
        return q("SELECT date, COUNT(*) AS bills, SUM(subtotal) AS subtotal, SUM(discount) AS discount, SUM(taxable) AS taxable, "
                 "SUM(cgst+sgst) AS tax, SUM(grand_total) AS total, "
                 "SUM(CASE WHEN payment_mode='Cash' THEN paid_amount ELSE 0 END) AS cash, "
                 "SUM(CASE WHEN payment_mode='UPI' THEN paid_amount ELSE 0 END) AS upi, "
                 "SUM(CASE WHEN payment_mode='Card' THEN paid_amount ELSE 0 END) AS card, "
                 "SUM(credit_amount) AS credit FROM sales WHERE branch_id=? AND date BETWEEN ? AND ? GROUP BY date ORDER BY date",
                 (bid, f, t))
    if kind == "gst":
        return q("SELECT si.gst_rate, SUM(si.taxable) AS taxable, SUM(si.tax)/2 AS cgst, SUM(si.tax)/2 AS sgst, SUM(si.total) AS total "
                 "FROM sale_items si JOIN sales s ON s.id=si.sale_id WHERE s.branch_id=? AND s.date BETWEEN ? AND ? "
                 "GROUP BY si.gst_rate ORDER BY si.gst_rate", (bid, f, t))
    if kind == "items":
        return q("SELECT si.name, si.unit, SUM(si.quantity) AS qty, SUM(si.total) AS sales, SUM(si.cost_price*si.quantity) AS cost, "
                 "SUM(si.total - si.cost_price*si.quantity) AS profit FROM sale_items si JOIN sales s ON s.id=si.sale_id "
                 "WHERE s.branch_id=? AND s.date BETWEEN ? AND ? GROUP BY si.product_id ORDER BY sales DESC", (bid, f, t))
    if kind == "counter":
        return q("SELECT co.code||' - '||co.name AS counter, u.full_name AS cashier, COUNT(*) AS bills, SUM(s.grand_total) AS total "
                 "FROM sales s JOIN counters co ON co.id=s.counter_id LEFT JOIN users u ON u.id=s.created_by "
                 "WHERE s.branch_id=? AND s.date BETWEEN ? AND ? GROUP BY s.counter_id, s.created_by ORDER BY co.code", (bid, f, t))
    if kind == "purchases":
        return q("SELECT p.date, p.invoice_no, s.name AS supplier, p.total FROM purchases p LEFT JOIN suppliers s ON s.id=p.supplier_id "
                 "WHERE p.branch_id=? AND p.date BETWEEN ? AND ? ORDER BY p.date", (bid, f, t))
    if kind == "branches":
        return q("SELECT b.name AS branch, COUNT(s.id) AS bills, COALESCE(SUM(s.grand_total),0) AS total FROM branches b "
                 "LEFT JOIN sales s ON s.branch_id=b.id AND s.date BETWEEN ? AND ? GROUP BY b.id ORDER BY b.name", (f, t))
    return []


REPORT_TITLES = {"daily": "Daily Sales Summary", "gst": "GST Summary", "items": "Item-wise Sales & Profit",
                 "counter": "Counter-wise Sales", "purchases": "Purchase Register", "branches": "Branch Comparison"}


@app.route("/reports")
@role_required("admin", "manager")
def reports():
    kind = request.args.get("kind", "daily")
    f = request.args.get("from") or date.today().replace(day=1).isoformat()
    t = request.args.get("to") or date.today().isoformat()
    rows = report_data(kind, f, t, branch_id())
    return render_template("reports.html", kind=kind, rows=rows, date_from=f, date_to=t, titles=REPORT_TITLES)


@app.route("/reports/export")
@role_required("admin", "manager")
def report_export():
    kind = request.args.get("kind", "daily")
    f, t = request.args.get("from"), request.args.get("to")
    rows = report_data(kind, f, t, branch_id())
    return export_response(REPORT_TITLES.get(kind, kind), f"{f} to {t}", rows, request.args.get("fmt", "xlsx"), filename=f"{kind}_report")


# --------------------------------------------------------------------------
# Settings, branches, counters, users
# --------------------------------------------------------------------------
@app.route("/settings", methods=["GET", "POST"])
@role_required("admin")
def settings_page():
    if request.method == "POST":
        sig = request.files.get("signature")
        if sig and sig.filename:
            ext = sig.filename.rsplit(".", 1)[-1].lower() if "." in sig.filename else ""
            if ext not in ("png", "jpg", "jpeg"):
                flash("Signature must be a PNG or JPG image.", "error")
                return redirect(url_for("settings_page"))
            os.makedirs(os.path.join(BASE_DIR, "uploads"), exist_ok=True)
            fname = f"signature.{ext}"
            sig.save(os.path.join(BASE_DIR, "uploads", fname))
            run("INSERT INTO settings(key,value) VALUES ('signature_file',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (fname,))
        if request.form.get("remove_signature"):
            run("UPDATE settings SET value='' WHERE key='signature_file'")
        for k in ["shop_name", "shop_tagline", "shop_address", "shop_phone", "shop_gstin", "state_name", "state_code", "invoice_footer",
                  "offer_letter_template", "emp_id_prefix", "designations", "id_card_validity_years", "signatory_name"]:
            run("INSERT INTO settings(key,value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (k, request.form.get(k, "").strip()))
        run("INSERT INTO settings(key,value) VALUES ('allow_negative_stock',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            ("1" if request.form.get("allow_negative_stock") else "0",))
        get_db().commit()
        flash("Settings saved.", "success")
        return redirect(url_for("settings_page"))
    return render_template("settings.html", sig_fonts=available_signature_fonts(), sig_inks=list(SIGNATURE_INKS))


SIGNATURE_FONTS = {                     # key: (label, Windows font file)
    "segoesc": ("Segoe Script", "segoesc.ttf"), "brush": ("Brush Script", "BRUSHSCI.TTF"), "freestyle": ("Freestyle Script", "FREESCPT.TTF"),
    "mistral": ("Mistral", "MISTRAL.TTF"), "corsiva": ("Monotype Corsiva", "MTCORSVA.TTF"), "pristina": ("Pristina", "PRISTINA.TTF"),
    "vivaldi": ("Vivaldi", "VIVALDII.TTF"), "lucida": ("Lucida Handwriting", "LHANDW.TTF"), "bradley": ("Bradley Hand", "BRADHITC.TTF"),
    "inkfree": ("Ink Free", "Inkfree.ttf"),
}
SIGNATURE_INKS = {"black": (20, 20, 20), "blue": (23, 37, 110), "navy": (12, 50, 90)}


def available_signature_fonts():
    out = {}
    for k, (label, f) in SIGNATURE_FONTS.items():
        path = os.path.join(r"C:\Windows\Fonts", f)
        if os.path.exists(path):
            out[k] = (label, path)
    return out


def render_signature(name, font_key="segoesc", ink="blue"):
    """Draw the typed name as a signature: transparent PNG with a light underline flourish."""
    from PIL import Image, ImageDraw, ImageFont
    fonts = available_signature_fonts()
    label, path = fonts.get(font_key) or next(iter(fonts.values()))
    size = 120
    font = ImageFont.truetype(path, size)
    name = name.strip() or "Signature"
    tmp = Image.new("RGBA", (10, 10))
    l, t, r, b = ImageDraw.Draw(tmp).textbbox((0, 0), name, font=font)
    w, h = r - l, b - t
    pad = 40
    img = Image.new("RGBA", (w + pad * 2, h + pad * 2 + 30), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    colour = SIGNATURE_INKS.get(ink, SIGNATURE_INKS["blue"]) + (255,)
    d.text((pad - l, pad - t), name, font=font, fill=colour)
    # flourish: a gentle curved underline
    y = pad + h + 14
    d.line([(pad + 8, y + 6), (pad + w * 0.35, y - 2), (pad + w * 0.7, y + 8), (pad + w + 4, y)], fill=colour, width=4, joint="curve")
    # slight slant to look hand written
    img = img.transform(img.size, Image.AFFINE, (1, 0.18, -0.18 * img.size[1] / 2, 0, 1, 0), resample=Image.BICUBIC)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


@app.route("/signature/preview")
@role_required("admin")
def signature_preview():
    png = render_signature(request.args.get("name", "Signature")[:40], request.args.get("font", "segoesc"), request.args.get("ink", "blue"))
    return send_file(io.BytesIO(png), mimetype="image/png")


@app.route("/settings/signature/generate", methods=["POST"])
@role_required("admin")
def signature_generate():
    name = request.form.get("name", "").strip()[:40]
    if not name:
        flash("Type the authorised person's name first.", "error")
        return redirect(url_for("settings_page"))
    png = render_signature(name, request.form.get("font", "segoesc"), request.form.get("ink", "blue"))
    os.makedirs(os.path.join(BASE_DIR, "uploads"), exist_ok=True)
    with open(os.path.join(BASE_DIR, "uploads", "signature.png"), "wb") as fh:
        fh.write(png)
    run("INSERT INTO settings(key,value) VALUES ('signature_file','signature.png') ON CONFLICT(key) DO UPDATE SET value=excluded.value")
    run("INSERT INTO settings(key,value) VALUES ('signatory_name',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (name,))
    get_db().commit()
    flash(f"Digital signature created for {name}. It now prints on all documents.", "success")
    return redirect(url_for("settings_page"))


@app.route("/signature")
@login_required
def signature_image():
    path = signature_path()
    if not path:
        abort(404)
    return send_from_directory(os.path.dirname(path), os.path.basename(path))


@app.route("/branches/<int:bid>/signature")
@login_required
def branch_signature(bid):
    path, _, _ = signature_for(bid)
    if not path:
        abort(404)
    return send_from_directory(os.path.dirname(path), os.path.basename(path))


@app.route("/branches/<int:bid>/signature", methods=["POST"])
@role_required("admin")
def branch_signature_save(bid):
    b = q("SELECT * FROM branches WHERE id=?", (bid,), one=True) or abort(404)
    f = request.form
    if f.get("remove"):
        if b["signature_file"]:
            try:
                os.remove(os.path.join(BASE_DIR, "uploads", b["signature_file"]))
            except OSError:
                pass
        run("UPDATE branches SET signature_file=NULL, signatory_name=NULL WHERE id=?", (bid,))
        get_db().commit()
        flash(f"{b['name']} will now use the company signature.", "success")
        return redirect(url_for("branches"))
    os.makedirs(os.path.join(BASE_DIR, "uploads"), exist_ok=True)
    file = request.files.get("signature")
    if file and file.filename:
        ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
        if ext not in ("png", "jpg", "jpeg"):
            flash("Signature must be a PNG or JPG image.", "error")
            return redirect(url_for("branches"))
        fname = f"signature_branch{bid}.{ext}"
        file.save(os.path.join(BASE_DIR, "uploads", fname))
    elif f.get("name", "").strip():
        fname = f"signature_branch{bid}.png"
        with open(os.path.join(BASE_DIR, "uploads", fname), "wb") as fh:
            fh.write(render_signature(f["name"].strip()[:40], f.get("font", "segoesc"), f.get("ink", "blue")))
    else:
        flash("Upload a signature image or type a name to create one.", "error")
        return redirect(url_for("branches"))
    run("UPDATE branches SET signature_file=?, signatory_name=? WHERE id=?", (fname, f.get("name", "").strip() or b["signatory_name"], bid))
    get_db().commit()
    flash(f"Signature saved for {b['name']}.", "success")
    return redirect(url_for("branches"))


@app.route("/branches", methods=["GET", "POST"])
@role_required("admin")
def branches():
    if request.method == "POST":
        f = request.form
        try:
            if f.get("id"):
                run("UPDATE branches SET code=?,name=?,address=?,phone=?,gstin=?,active=? WHERE id=?",
                    (f["code"].strip().upper(), f["name"].strip(), f.get("address", ""), f.get("phone", ""), f.get("gstin", ""),
                     1 if f.get("active") else 0, f["id"]))
            else:
                bid = run("INSERT INTO branches(code,name,address,phone,gstin) VALUES (?,?,?,?,?)",
                          (f["code"].strip().upper(), f["name"].strip(), f.get("address", ""), f.get("phone", ""), f.get("gstin", "")))
                run("INSERT INTO counters(branch_id,code,name) VALUES (?,'C1','Counter 1')", (bid,))
            get_db().commit()
        except sqlite3.IntegrityError:
            flash("Branch code must be unique.", "error")
        return redirect(url_for("branches"))
    rows = q("SELECT b.*, (SELECT COUNT(*) FROM counters c WHERE c.branch_id=b.id AND c.active=1) AS n_counters FROM branches b ORDER BY b.name")
    counters = q("SELECT c.*, b.name AS branch_name FROM counters c JOIN branches b ON b.id=c.branch_id ORDER BY b.name, c.code")
    return render_template("branches.html", branches=rows, counters=counters, sig_fonts=available_signature_fonts(), sig_inks=list(SIGNATURE_INKS))


@app.route("/counters", methods=["POST"])
@role_required("admin")
def counters_save():
    f = request.form
    try:
        if f.get("id"):
            run("UPDATE counters SET code=?,name=?,active=? WHERE id=?",
                (f["code"].strip().upper(), f["name"].strip(), 1 if f.get("active") else 0, f["id"]))
        else:
            run("INSERT INTO counters(branch_id,code,name) VALUES (?,?,?)",
                (f["branch_id"], f["code"].strip().upper(), f["name"].strip()))
        get_db().commit()
    except sqlite3.IntegrityError:
        flash("Counter code must be unique within the branch.", "error")
    return redirect(url_for("branches"))


@app.route("/users", methods=["GET", "POST"])
@login_required
def users():
    u = current_user()
    is_admin = u["role"] == "admin"
    if not (is_admin or (u["role"] == "manager" and u["can_manage_users"])):
        flash("You do not have permission for that page.", "error")
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        f = request.form
        if is_admin:
            role = f["role"]
            bid = None if role == "admin" else (f.get("branch_id") or None)
            can_manage = 1 if (role == "manager" and f.get("can_manage_users")) else 0
            can_analytics = 1 if (role != "admin" and f.get("can_view_analytics")) else 0
            device_lock = 1 if f.get("device_lock") else 0
            can_reset = 1 if (role == "manager" and f.get("can_reset_device")) else 0
            can_hr = 1 if (role == "manager" and f.get("can_hr_docs")) else 0
        else:                                   # manager: cashiers of own branch only
            role, bid, can_manage, can_analytics = "cashier", u["branch_id"], 0, 0
            device_lock = 1 if f.get("device_lock", "1") else 0
            can_reset = can_hr = 0
        joining = f.get("joining_date") or None
        try:
            if f.get("id"):
                target = q("SELECT * FROM users WHERE id=?", (f["id"],), one=True)
                if not target or (not is_admin and (target["role"] != "cashier" or target["branch_id"] != u["branch_id"])):
                    flash("You can only edit cashiers of your own branch.", "error")
                    return redirect(url_for("users"))
                run("UPDATE users SET full_name=?,role=?,branch_id=?,active=?,joining_date=?,can_manage_users=?,can_view_analytics=?,device_lock=?,"
                    "can_reset_device=?,can_hr_docs=? WHERE id=?",
                    (f["full_name"].strip(), role, bid, 1 if f.get("active") else 0, joining, can_manage, can_analytics, device_lock, can_reset, can_hr, f["id"]))
                if f.get("designation") is not None:
                    run("INSERT INTO staff_profiles(user_id,emp_code,designation) VALUES (?,?,?) "
                        "ON CONFLICT(user_id) DO UPDATE SET designation=excluded.designation", (f["id"], next_emp_code(), f.get("designation", "").strip()))
                if not device_lock:
                    run("UPDATE users SET bound_ip=NULL, bound_at=NULL WHERE id=?", (f["id"],))   # lock off = forget the device
                if f.get("password"):
                    run("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(f["password"]), f["id"]))
            else:
                new_id = run("INSERT INTO users(username,password_hash,full_name,role,branch_id,joining_date,can_manage_users,can_view_analytics,device_lock,"
                             "can_reset_device,can_hr_docs) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                             (f["username"].strip().lower(), generate_password_hash(f["password"]), f["full_name"].strip(), role, bid, joining, can_manage,
                              can_analytics, device_lock, can_reset, can_hr))
                run("INSERT INTO staff_profiles(user_id,emp_code,designation) VALUES (?,?,?)",
                    (new_id, next_emp_code(), f.get("designation", "").strip() or role.title()))
            get_db().commit()
            flash("User saved.", "success")
        except sqlite3.IntegrityError:
            flash("Username already exists.", "error")
        return redirect(url_for("users"))
    if is_admin:
        rows = q("SELECT u.*, b.name AS branch_name, p.emp_code, p.designation FROM users u LEFT JOIN branches b ON b.id=u.branch_id "
                 "LEFT JOIN staff_profiles p ON p.user_id=u.id ORDER BY u.username")
    else:
        rows = q("SELECT u.*, b.name AS branch_name, p.emp_code, p.designation FROM users u LEFT JOIN branches b ON b.id=u.branch_id "
                 "LEFT JOIN staff_profiles p ON p.user_id=u.id WHERE u.role='cashier' AND u.branch_id=? ORDER BY u.username", (u["branch_id"],))
    return render_template("users.html", users=rows, branches=q("SELECT * FROM branches WHERE active=1 ORDER BY name"), is_admin=is_admin,
                           designations=designation_list(), next_code=next_emp_code())


@app.route("/users/<int:uid>/reset-device", methods=["POST"])
@login_required
def user_reset_device(uid):
    u = current_user()
    target = q("SELECT * FROM users WHERE id=?", (uid,), one=True) or abort(404)
    allowed = u["role"] == "admin" or (u["role"] == "manager" and u["can_reset_device"] and target["role"] == "cashier"
                                       and target["branch_id"] == u["branch_id"])
    if not allowed:
        abort(403)
    run("UPDATE users SET bound_ip=NULL, bound_at=NULL WHERE id=?", (uid,))
    get_db().commit()
    flash(f"Device reset for {target['full_name']}. The next login will register the new device.", "success")
    return redirect(request.referrer or url_for("devices"))


@app.route("/devices")
@login_required
def devices():
    u = current_user()
    if u["role"] == "admin":
        rows = q("SELECT u.*, b.name AS branch_name FROM users u LEFT JOIN branches b ON b.id=u.branch_id WHERE u.active=1 ORDER BY b.name, u.full_name")
    elif u["role"] == "manager" and u["can_reset_device"]:
        rows = q("SELECT u.*, b.name AS branch_name FROM users u LEFT JOIN branches b ON b.id=u.branch_id "
                 "WHERE u.active=1 AND u.role='cashier' AND u.branch_id=? ORDER BY u.full_name", (u["branch_id"],))
    else:
        flash("You do not have permission for that page.", "error")
        return redirect(url_for("dashboard"))
    blocked = q("SELECT a.at, a.username, a.ip, a.details FROM audit_log a WHERE a.action='Login BLOCKED (device)' ORDER BY a.id DESC LIMIT 20")
    return render_template("devices.html", rows=rows, blocked=blocked)


@app.route("/change-password", methods=["GET", "POST"])
@login_required
def change_password():
    if request.method == "POST":
        u = current_user()
        if not check_password_hash(u["password_hash"], request.form["old"]):
            flash("Current password is wrong.", "error")
        elif request.form["new"] != request.form["confirm"] or len(request.form["new"]) < 4:
            flash("New passwords do not match or are too short.", "error")
        else:
            run("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(request.form["new"]), u["id"]))
            get_db().commit()
            flash("Password changed.", "success")
            return redirect(url_for("dashboard"))
    return render_template("change_password.html")


# --------------------------------------------------------------------------
# Messages between branches (and stock requests)
# --------------------------------------------------------------------------
def transfer_stock(product_id, from_b, to_b, qty, reason):
    """Move stock between branches and log both sides."""
    t = date.today().isoformat()
    change_stock(product_id, from_b, -qty)
    change_stock(product_id, to_b, qty)
    run("INSERT INTO stock_adjustments(branch_id,product_id,change_qty,reason,date,created_by) VALUES (?,?,?,?,?,?)",
        (from_b, product_id, -qty, reason, t, session["user_id"]))
    run("INSERT INTO stock_adjustments(branch_id,product_id,change_qty,reason,date,created_by) VALUES (?,?,?,?,?,?)",
        (to_b, product_id, qty, reason, t, session["user_id"]))


@app.route("/messages")
@role_required("admin", "manager")
def messages():
    b = branch_id()
    box = request.args.get("box", "inbox")
    base = ("SELECT m.*, fb.name AS from_branch, tb.name AS to_branch, u.full_name AS sender, p.name AS product, p.unit, "
            "(SELECT COUNT(*) FROM messages r WHERE r.parent_id=m.id) AS replies, "
            "(SELECT COUNT(*) FROM messages r WHERE (r.id=m.id OR r.parent_id=m.id) AND r.to_branch_id=? AND r.is_read=0) AS unread, "
            "(SELECT MAX(created_at) FROM messages r WHERE r.id=m.id OR r.parent_id=m.id) AS last_at "
            "FROM messages m JOIN branches fb ON fb.id=m.from_branch_id JOIN branches tb ON tb.id=m.to_branch_id "
            "LEFT JOIN users u ON u.id=m.from_user_id LEFT JOIN products p ON p.id=m.product_id WHERE m.parent_id IS NULL ")
    if box == "sent":
        rows = q(base + "AND m.from_branch_id=? ORDER BY last_at DESC", (b, b))
    elif box == "requests":
        rows = q(base + "AND m.kind='stock_request' AND (m.to_branch_id=? OR m.from_branch_id=?) ORDER BY "
                 "CASE m.status WHEN 'open' THEN 0 ELSE 1 END, last_at DESC", (b, b, b))
    else:
        rows = q(base + "AND m.to_branch_id=? ORDER BY unread DESC, last_at DESC", (b, b))
    return render_template("messages.html", rows=rows, box=box)


@app.route("/messages/new", methods=["GET", "POST"])
@role_required("admin", "manager")
def message_new():
    b = branch_id()
    others = q("SELECT * FROM branches WHERE active=1 AND id<>? ORDER BY name", (b,))
    if request.method == "POST":
        f = request.form
        kind = f.get("kind", "message")
        to_b = int(f["to_branch_id"])
        if to_b == b or not any(o["id"] == to_b for o in others):
            flash("Choose another branch.", "error")
            return redirect(url_for("message_new"))
        product_id = qty = None
        subject = f.get("subject", "").strip()
        if kind == "stock_request":
            product_id = int(f.get("product_id") or 0) or None
            qty = float(f.get("quantity") or 0)
            if not product_id or qty <= 0:
                flash("Pick a product and a quantity for a stock request.", "error")
                return redirect(url_for("message_new", kind="stock_request"))
            pr = q("SELECT name, unit FROM products WHERE id=?", (product_id,), one=True)
            subject = subject or f"Stock request: {qty_fmt(qty)} {pr['unit']} {pr['name']}"
        if not subject and not f.get("body", "").strip():
            flash("Write a subject or a message.", "error")
            return redirect(url_for("message_new"))
        mid = run("INSERT INTO messages(from_user_id,from_branch_id,to_branch_id,kind,subject,body,product_id,quantity) "
                  "VALUES (?,?,?,?,?,?,?,?)",
                  (session["user_id"], b, to_b, kind, subject, f.get("body", "").strip(), product_id, qty))
        get_db().commit()
        flash("Message sent.", "success")
        return redirect(url_for("message_thread", mid=mid))
    return render_template("message_new.html", others=others, kind=request.args.get("kind", "message"),
                           product_id=request.args.get("product_id", ""), product_name=request.args.get("product_name", ""))


@app.route("/api/branch-stock")
@login_required
def api_branch_stock():
    """Stock of one product at every branch (used by the stock request form)."""
    pid = request.args.get("product_id", type=int)
    rows = q("SELECT b.id, b.name, COALESCE(s.quantity,0) AS qty FROM branches b "
             "LEFT JOIN stock s ON s.branch_id=b.id AND s.product_id=? WHERE b.active=1 ORDER BY b.name", (pid,))
    return jsonify([dict(r) for r in rows])


def load_thread(mid):
    root = q("SELECT m.*, fb.name AS from_branch, tb.name AS to_branch, u.full_name AS sender, p.name AS product, p.unit "
             "FROM messages m JOIN branches fb ON fb.id=m.from_branch_id JOIN branches tb ON tb.id=m.to_branch_id "
             "LEFT JOIN users u ON u.id=m.from_user_id LEFT JOIN products p ON p.id=m.product_id "
             "WHERE m.id=? AND m.parent_id IS NULL", (mid,), one=True)
    if not root:
        abort(404)
    u = current_user()
    if u["role"] != "admin" and branch_id() not in (root["from_branch_id"], root["to_branch_id"]):
        abort(403)
    return root


@app.route("/messages/<int:mid>", methods=["GET", "POST"])
@role_required("admin", "manager")
def message_thread(mid):
    root = load_thread(mid)
    b = branch_id()
    other_b = root["from_branch_id"] if b == root["to_branch_id"] else root["to_branch_id"]
    if request.method == "POST":
        body = request.form.get("body", "").strip()
        if body:
            run("INSERT INTO messages(parent_id,from_user_id,from_branch_id,to_branch_id,kind,body) VALUES (?,?,?,?,'message',?)",
                (mid, session["user_id"], b, other_b, body))
            get_db().commit()
        return redirect(url_for("message_thread", mid=mid))
    run("UPDATE messages SET is_read=1 WHERE (id=? OR parent_id=?) AND to_branch_id=?", (mid, mid, b))
    get_db().commit()
    replies = q("SELECT m.*, fb.name AS from_branch, u.full_name AS sender FROM messages m JOIN branches fb ON fb.id=m.from_branch_id "
                "LEFT JOIN users u ON u.id=m.from_user_id WHERE m.parent_id=? ORDER BY m.id", (mid,))
    my_stock = their_stock = None
    if root["kind"] == "stock_request" and root["product_id"]:
        st = {r["id"]: r["qty"] for r in q("SELECT branch_id AS id, quantity AS qty FROM stock WHERE product_id=?", (root["product_id"],))}
        my_stock, their_stock = st.get(b, 0), st.get(other_b, 0)
    return render_template("message_thread.html", root=root, replies=replies, my_stock=my_stock, their_stock=their_stock,
                           i_am_receiver=(b == root["to_branch_id"]), i_am_sender=(b == root["from_branch_id"]))


@app.route("/messages/<int:mid>/send-stock", methods=["POST"])
@role_required("admin", "manager")
def message_send_stock(mid):
    """The receiving branch fulfils a stock request: stock moves from this branch to the requesting branch."""
    root = load_thread(mid)
    b = branch_id()
    if root["kind"] != "stock_request" or b != root["to_branch_id"] or root["status"] != "open":
        flash("This request cannot be fulfilled from here.", "error")
        return redirect(url_for("message_thread", mid=mid))
    qty = float(request.form.get("quantity") or 0)
    have = q("SELECT COALESCE(quantity,0) AS qty FROM stock WHERE product_id=? AND branch_id=?", (root["product_id"], b), one=True)
    have = have["qty"] if have else 0
    if qty <= 0 or qty > have + 0.0005:
        flash(f"You can send at most {qty_fmt(have)} {root['unit']}.", "error")
        return redirect(url_for("message_thread", mid=mid))
    transfer_stock(root["product_id"], b, root["from_branch_id"], qty, f"Stock request #{mid}: {root['to_branch']} -> {root['from_branch']}")
    run("UPDATE messages SET status='sent', sent_quantity=? WHERE id=?", (qty, mid))
    note = request.form.get("body", "").strip()
    run("INSERT INTO messages(parent_id,from_user_id,from_branch_id,to_branch_id,kind,body) VALUES (?,?,?,?,'message',?)",
        (mid, session["user_id"], b, root["from_branch_id"],
         f"Sent {qty_fmt(qty)} {root['unit']} of {root['product']}. Stock has been transferred." + (f" {note}" if note else "")))
    get_db().commit()
    flash(f"Sent {qty_fmt(qty)} {root['unit']} to {root['from_branch']}. Stock updated on both branches.", "success")
    return redirect(url_for("message_thread", mid=mid))


@app.route("/messages/<int:mid>/status", methods=["POST"])
@role_required("admin", "manager")
def message_status(mid):
    root = load_thread(mid)
    b = branch_id()
    new = request.form.get("status")
    if new == "declined" and b == root["to_branch_id"] and root["status"] == "open":
        pass
    elif new == "closed" and (b == root["from_branch_id"] or current_user()["role"] == "admin"):
        pass
    elif new == "open" and root["status"] == "declined" and b == root["from_branch_id"]:
        pass
    else:
        flash("Not allowed.", "error")
        return redirect(url_for("message_thread", mid=mid))
    run("UPDATE messages SET status=? WHERE id=?", (new, mid))
    note = request.form.get("body", "").strip()
    text = {"declined": "Request declined.", "closed": "Thread closed.", "open": "Request reopened."}[new]
    other_b = root["from_branch_id"] if b == root["to_branch_id"] else root["to_branch_id"]
    run("INSERT INTO messages(parent_id,from_user_id,from_branch_id,to_branch_id,kind,body) VALUES (?,?,?,?,'message',?)",
        (mid, session["user_id"], b, other_b, text + (f" {note}" if note else "")))
    get_db().commit()
    return redirect(url_for("message_thread", mid=mid))


# --------------------------------------------------------------------------
# Notices (admin -> branches / staff / managers, manager -> own branch staff)
# --------------------------------------------------------------------------
@app.route("/notices")
@login_required
def notices():
    u = current_user()
    box = request.args.get("box", "inbox")
    if box == "sent" and u["role"] in ("admin", "manager"):
        sql = ("SELECT n.*, u.full_name AS sender, (SELECT COUNT(*) FROM notice_recipients r WHERE r.notice_id=n.id) AS total, "
               "(SELECT COUNT(*) FROM notice_recipients r WHERE r.notice_id=n.id AND r.read_at IS NOT NULL) AS read_count "
               "FROM notices n JOIN users u ON u.id=n.from_user_id ")
        rows = q(sql + "ORDER BY n.id DESC") if u["role"] == "admin" else q(sql + "WHERE n.from_user_id=? ORDER BY n.id DESC", (u["id"],))
    else:
        box = "inbox"
        rows = q("SELECT n.*, u.full_name AS sender, r.read_at FROM notice_recipients r JOIN notices n ON n.id=r.notice_id "
                 "JOIN users u ON u.id=n.from_user_id WHERE r.user_id=? ORDER BY r.read_at IS NOT NULL, n.id DESC", (u["id"],))
    return render_template("notices.html", rows=rows, box=box)


@app.route("/notices/new", methods=["GET", "POST"])
@role_required("admin", "manager")
def notice_new():
    u = current_user()
    is_admin = u["role"] == "admin"
    if is_admin:
        staff = q("SELECT u.*, b.name AS branch_name FROM users u LEFT JOIN branches b ON b.id=u.branch_id "
                  "WHERE u.active=1 AND u.id<>? ORDER BY b.name, u.full_name", (u["id"],))
    else:
        staff = q("SELECT u.*, b.name AS branch_name FROM users u LEFT JOIN branches b ON b.id=u.branch_id "
                  "WHERE u.active=1 AND u.id<>? AND u.branch_id=? ORDER BY u.full_name", (u["id"], u["branch_id"]))
    branches_list = q("SELECT * FROM branches WHERE active=1 ORDER BY name")
    if request.method == "POST":
        f = request.form
        title, body = f.get("title", "").strip(), f.get("body", "").strip()
        if not title:
            flash("Give the notice a title.", "error")
            return redirect(url_for("notice_new"))
        scope = f.get("scope", "all")
        managers_only = bool(f.get("managers_only")) and is_admin
        chosen = staff
        if scope == "users":
            ids = {int(x) for x in f.getlist("user_ids")}
            chosen = [x for x in staff if x["id"] in ids]
            audience = f"Selected staff ({len(chosen)})"
        elif scope == "branches" and is_admin:
            bids = {int(x) for x in f.getlist("branch_ids")}
            chosen = [x for x in staff if x["branch_id"] in bids]
            names = ", ".join(b["name"] for b in branches_list if b["id"] in bids)
            audience = f"{'Managers' if managers_only else 'All staff'} of {names or 'no branch'}"
        else:
            if is_admin:
                audience = "All managers, all branches" if managers_only else "All staff, all branches"
            else:
                bname = q("SELECT name FROM branches WHERE id=?", (u["branch_id"],), one=True)["name"]
                audience = f"All staff of {bname}"
        if managers_only and scope != "users":
            chosen = [x for x in chosen if x["role"] == "manager"]
        if not chosen:
            flash("No staff matched that selection.", "error")
            return redirect(url_for("notice_new"))
        nid = run("INSERT INTO notices(from_user_id,title,body,audience) VALUES (?,?,?,?)", (u["id"], title, body, audience))
        for x in chosen:
            run("INSERT OR IGNORE INTO notice_recipients(notice_id,user_id) VALUES (?,?)", (nid, x["id"]))
        get_db().commit()
        flash(f"Notice sent to {len(chosen)} staff member(s).", "success")
        return redirect(url_for("notice_detail", nid=nid))
    return render_template("notice_new.html", staff=staff, branches_list=branches_list, is_admin=is_admin)


@app.route("/notices/<int:nid>")
@login_required
def notice_detail(nid):
    u = current_user()
    n = q("SELECT n.*, u.full_name AS sender, u.role AS sender_role FROM notices n JOIN users u ON u.id=n.from_user_id WHERE n.id=?",
          (nid,), one=True) or abort(404)
    rcpt = q("SELECT * FROM notice_recipients WHERE notice_id=? AND user_id=?", (nid, u["id"]), one=True)
    is_sender = n["from_user_id"] == u["id"] or u["role"] == "admin"
    if not rcpt and not is_sender:
        abort(403)
    if rcpt and not rcpt["read_at"]:
        run("UPDATE notice_recipients SET read_at=datetime('now','localtime') WHERE notice_id=? AND user_id=?", (nid, u["id"]))
        get_db().commit()
    recipients = q("SELECT r.read_at, u.full_name, u.role, b.name AS branch_name FROM notice_recipients r JOIN users u ON u.id=r.user_id "
                   "LEFT JOIN branches b ON b.id=u.branch_id WHERE r.notice_id=? ORDER BY r.read_at IS NOT NULL, b.name, u.full_name",
                   (nid,)) if is_sender else []
    return render_template("notice_detail.html", n=n, recipients=recipients, is_sender=is_sender)


@app.route("/notices/<int:nid>/delete", methods=["POST"])
@role_required("admin", "manager")
def notice_delete(nid):
    u = current_user()
    n = q("SELECT * FROM notices WHERE id=?", (nid,), one=True) or abort(404)
    if u["role"] != "admin" and n["from_user_id"] != u["id"]:
        abort(403)
    run("DELETE FROM notice_recipients WHERE notice_id=?", (nid,))
    run("DELETE FROM notices WHERE id=?", (nid,))
    get_db().commit()
    flash("Notice deleted.", "success")
    return redirect(url_for("notices", box="sent"))


@app.route("/api/notices/unread")
@login_required
def api_notices_unread():
    """Polled by every page so a new notice pops up without reloading."""
    return jsonify([dict(r) for r in unread_notice_rows(session["user_id"])])


@app.route("/notices/<int:nid>/read", methods=["POST"])
@login_required
def notice_mark_read(nid):
    run("UPDATE notice_recipients SET read_at=datetime('now','localtime') WHERE notice_id=? AND user_id=? AND read_at IS NULL",
        (nid, session["user_id"]))
    get_db().commit()
    left = q("SELECT COUNT(*) AS n FROM notice_recipients WHERE user_id=? AND read_at IS NULL", (session["user_id"],), one=True)["n"]
    return jsonify(ok=True, unread=left)

# --------------------------------------------------------------------------
# Leave applications
# --------------------------------------------------------------------------
LEAVE_SQL = ("SELECT l.*, u.full_name AS applicant, u.role AS applicant_role, b.name AS branch_name, "
             "d.full_name AS decided_by_name, o.full_name AS overridden_by_name FROM leaves l JOIN users u ON u.id=l.user_id "
             "LEFT JOIN branches b ON b.id=l.branch_id LEFT JOIN users d ON d.id=l.decided_by LEFT JOIN users o ON o.id=l.overridden_by ")


def can_decide(u, leave):
    """Who may approve/reject: branch managers for cashiers of their branch, admin for everyone (incl. managers)."""
    if u["id"] == leave["user_id"]:
        return False
    if u["role"] == "admin":
        return True
    return u["role"] == "manager" and leave["applicant_role"] == "cashier" and leave["branch_id"] == u["branch_id"]


def pending_leaves_for(u):
    if u["role"] == "admin":
        return q(LEAVE_SQL + "WHERE l.status='pending' AND l.user_id<>? ORDER BY l.id", (u["id"],))
    if u["role"] == "manager":
        return q(LEAVE_SQL + "WHERE l.status='pending' AND u.role='cashier' AND l.branch_id=? ORDER BY l.id", (u["branch_id"],))
    return []


def notify(user_ids, title, body, from_user_id):
    """Create a notice (so the popup shows it) for the given users."""
    user_ids = [x for x in set(user_ids) if x and x != from_user_id]
    if not user_ids:
        return
    nid = run("INSERT INTO notices(from_user_id,title,body,audience) VALUES (?,?,?,?)", (from_user_id, title, body, "Leave notification"))
    for uid in user_ids:
        run("INSERT OR IGNORE INTO notice_recipients(notice_id,user_id) VALUES (?,?)", (nid, uid))


def approvers_for(u):
    if u["role"] == "manager" or u["role"] == "admin":
        return [r["id"] for r in q("SELECT id FROM users WHERE role='admin' AND active=1")]
    return [r["id"] for r in q("SELECT id FROM users WHERE role='manager' AND active=1 AND branch_id=?", (u["branch_id"],))] or \
           [r["id"] for r in q("SELECT id FROM users WHERE role='admin' AND active=1")]


def save_document(file):
    if not file or not file.filename:
        return None, None
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_DOCS:
        raise ValueError("Only PDF, JPG and PNG files are allowed.")
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    stored = f"{session['user_id']}_{int(datetime.now().timestamp())}_{secure_filename(file.filename)[:60]}"
    file.save(os.path.join(UPLOAD_DIR, stored))
    return stored, file.filename


@app.route("/leaves")
@login_required
def leaves():
    u = current_user()
    box = request.args.get("box", "mine")
    if box == "approvals" and u["role"] != "cashier":
        rows = pending_leaves_for(u)
    elif box == "all" and u["role"] != "cashier":
        rows = q(LEAVE_SQL + "ORDER BY l.id DESC LIMIT 300") if u["role"] == "admin" else \
            q(LEAVE_SQL + "WHERE l.branch_id=? ORDER BY l.id DESC LIMIT 300", (u["branch_id"],))
    else:
        box = "mine"
        rows = q(LEAVE_SQL + "WHERE l.user_id=? ORDER BY l.id DESC", (u["id"],))
    return render_template("leaves.html", rows=rows, box=box)


@app.route("/leave-policy", methods=["GET", "POST"])
@role_required("admin")
def leave_policy():
    if request.method == "POST":
        f = request.form
        if f.get("new_type"):
            t = f["new_type"].strip().title()
            try:
                run("INSERT INTO leave_policies(leave_type,days_per_year,eligible_after_months,sort_order) VALUES (?,?,?,?)",
                    (t, int(f.get("new_days") or 0), int(f.get("new_months") or 0),
                     q("SELECT COALESCE(MAX(sort_order),0)+1 AS n FROM leave_policies", one=True)["n"]))
            except sqlite3.IntegrityError:
                flash("That leave type already exists.", "error")
        else:
            for t in f.getlist("type"):
                run("UPDATE leave_policies SET days_per_year=?, eligible_after_months=?, active=? WHERE leave_type=?",
                    (int(f.get(f"days_{t}") or 0), int(f.get(f"months_{t}") or 0), 1 if f.get(f"active_{t}") else 0, t))
        get_db().commit()
        flash("Leave policy saved.", "success")
        return redirect(url_for("leave_policy"))
    return render_template("leave_policy.html", policies=q("SELECT * FROM leave_policies ORDER BY sort_order, leave_type"))


def add_months(d, months):
    m = d.month - 1 + months
    y, m = d.year + m // 12, m % 12 + 1
    import calendar
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def leave_balances(user, year=None):
    """Per active leave type: entitlement, eligibility date, used (approved) and pending days in the year."""
    year = year or date.today().year
    joined = date.fromisoformat(user["joining_date"]) if user["joining_date"] else None
    out = []
    for pol in q("SELECT * FROM leave_policies WHERE active=1 ORDER BY sort_order, leave_type"):
        t = pol["leave_type"]
        used = q("SELECT COALESCE(SUM(days),0) AS d FROM leaves WHERE user_id=? AND leave_type=? AND status='approved' AND substr(from_date,1,4)=?",
                 (user["id"], t, str(year)), one=True)["d"]
        pend = q("SELECT COALESCE(SUM(days),0) AS d FROM leaves WHERE user_id=? AND leave_type=? AND status='pending' AND substr(from_date,1,4)=?",
                 (user["id"], t, str(year)), one=True)["d"]
        eligible_from = add_months(joined, pol["eligible_after_months"]) if (joined and pol["eligible_after_months"]) else joined
        limited = pol["days_per_year"] > 0
        out.append({"type": t, "entitled": pol["days_per_year"], "limited": limited, "months": pol["eligible_after_months"],
                    "eligible_from": eligible_from, "used": used, "pending": pend,
                    "balance": (pol["days_per_year"] - used - pend) if limited else None})
    return out


@app.route("/leaves/new", methods=["GET", "POST"])
@login_required
def leave_new():
    u = current_user()
    balances = leave_balances(u)
    if request.method == "POST":
        f = request.form
        try:
            d1, d2 = date.fromisoformat(f["from_date"]), date.fromisoformat(f["to_date"])
        except (ValueError, KeyError):
            flash("Enter valid dates.", "error")
            return redirect(url_for("leave_new"))
        if d2 < d1:
            flash("End date cannot be before start date.", "error")
            return redirect(url_for("leave_new"))
        ltype = f.get("leave_type")
        bal = next((b for b in balances if b["type"] == ltype), None)
        if not bal:
            flash("Choose a valid leave type.", "error")
            return redirect(url_for("leave_new"))
        if bal["eligible_from"] and d1 < bal["eligible_from"]:
            flash(f"{ltype} leave is available to you only from {bal['eligible_from'].isoformat()} "
                  f"({bal['months']} month(s) after joining).", "error")
            return redirect(url_for("leave_new"))
        days = (d2 - d1).days + 1
        if bal["limited"] and days > bal["balance"]:
            flash(f"Only {bal['balance']} day(s) of {ltype} leave left this year (entitled {bal['entitled']}, "
                  f"used {bal['used']}, pending {bal['pending']}). Apply for fewer days or choose another type.", "error")
            return redirect(url_for("leave_new"))
        try:
            stored, orig = save_document(request.files.get("document"))
        except ValueError as e:
            flash(str(e), "error")
            return redirect(url_for("leave_new"))
        lid = run("INSERT INTO leaves(user_id,branch_id,leave_type,from_date,to_date,days,reason,document,document_name) "
                  "VALUES (?,?,?,?,?,?,?,?,?)",
                  (u["id"], u["branch_id"] or session.get("branch_id"), ltype, d1.isoformat(), d2.isoformat(), (d2 - d1).days + 1,
                   f.get("reason", "").strip(), stored, orig))
        notify(approvers_for(u), f"Leave application from {u['full_name']}",
               f"{ltype} leave {d1.isoformat()} to {d2.isoformat()} ({(d2 - d1).days + 1} day(s)). Reason: {f.get('reason', '').strip() or '-'}\n"
               f"Open Leave > Pending approvals to approve or reject.", u["id"])
        get_db().commit()
        flash("Leave application submitted.", "success")
        return redirect(url_for("leave_detail", lid=lid))
    return render_template("leave_new.html", balances=balances)


@app.route("/leaves/<int:lid>")
@login_required
def leave_detail(lid):
    u = current_user()
    l = q(LEAVE_SQL + "WHERE l.id=?", (lid,), one=True) or abort(404)
    is_owner = l["user_id"] == u["id"]
    if not (is_owner or u["role"] == "admin" or (u["role"] == "manager" and l["branch_id"] == u["branch_id"])):
        abort(403)
    applicant = q("SELECT * FROM users WHERE id=?", (l["user_id"],), one=True)
    bal = next((b for b in leave_balances(applicant, int(l["from_date"][:4])) if b["type"] == l["leave_type"]), None)
    return render_template("leave_detail.html", l=l, is_owner=is_owner, bal=bal, can_decide=can_decide(u, l) and l["status"] == "pending",
                           can_override=(u["role"] == "admin" and l["status"] in ("approved", "rejected") and not is_owner))


@app.route("/leaves/<int:lid>/decide", methods=["POST"])
@login_required
def leave_decide(lid):
    u = current_user()
    l = q(LEAVE_SQL + "WHERE l.id=?", (lid,), one=True) or abort(404)
    new = request.form.get("status")
    note = request.form.get("note", "").strip()
    if new not in ("approved", "rejected"):
        abort(400)
    if l["status"] == "pending" and can_decide(u, l):
        run("UPDATE leaves SET status=?, decided_by=?, decided_at=datetime('now','localtime'), decision_note=? WHERE id=?",
            (new, u["id"], note, lid))
        notify([l["user_id"]], f"Leave {new}: {l['from_date']} to {l['to_date']}",
               f"Your {l['leave_type']} leave was {new} by {u['full_name']}." + (f" Note: {note}" if note else ""), u["id"])
        flash(f"Leave {new}.", "success")
    elif u["role"] == "admin" and l["status"] in ("approved", "rejected") and l["user_id"] != u["id"]:
        if new == l["status"]:
            flash("The leave already has that status.", "error")
            return redirect(url_for("leave_detail", lid=lid))
        run("UPDATE leaves SET status=?, overridden_by=?, overridden_at=datetime('now','localtime'), override_note=? WHERE id=?",
            (new, u["id"], note, lid))
        notify([l["user_id"], l["decided_by"]], f"Leave decision changed to {new}: {l['from_date']} to {l['to_date']}",
               f"Admin {u['full_name']} changed the decision on {l['applicant']}'s {l['leave_type']} leave from {l['status']} to {new}."
               + (f" Note: {note}" if note else ""), u["id"])
        flash(f"Decision overridden: now {new}.", "success")
    else:
        flash("You cannot decide this application.", "error")
    get_db().commit()
    return redirect(url_for("leave_detail", lid=lid))


@app.route("/leaves/<int:lid>/cancel", methods=["POST"])
@login_required
def leave_cancel(lid):
    l = q("SELECT * FROM leaves WHERE id=?", (lid,), one=True) or abort(404)
    if l["user_id"] != session["user_id"] or l["status"] != "pending":
        flash("Only a pending application can be withdrawn by the applicant.", "error")
    else:
        run("UPDATE leaves SET status='cancelled' WHERE id=?", (lid,))
        get_db().commit()
        flash("Application withdrawn.", "success")
    return redirect(url_for("leave_detail", lid=lid))


@app.route("/leaves/<int:lid>/document", methods=["GET", "POST"])
@login_required
def leave_document(lid):
    u = current_user()
    l = q(LEAVE_SQL + "WHERE l.id=?", (lid,), one=True) or abort(404)
    is_owner = l["user_id"] == u["id"]
    if not (is_owner or u["role"] == "admin" or (u["role"] == "manager" and l["branch_id"] == u["branch_id"])):
        abort(403)
    if request.method == "POST":
        if not is_owner:
            abort(403)
        try:
            stored, orig = save_document(request.files.get("document"))
        except ValueError as e:
            flash(str(e), "error")
            return redirect(url_for("leave_detail", lid=lid))
        if not stored:
            flash("Choose a file first.", "error")
            return redirect(url_for("leave_detail", lid=lid))
        if l["document"]:
            try:
                os.remove(os.path.join(UPLOAD_DIR, l["document"]))
            except OSError:
                pass
        run("UPDATE leaves SET document=?, document_name=? WHERE id=?", (stored, orig, lid))
        get_db().commit()
        flash("Document uploaded.", "success")
        return redirect(url_for("leave_detail", lid=lid))
    if not l["document"]:
        abort(404)
    return send_from_directory(UPLOAD_DIR, l["document"], as_attachment=False, download_name=l["document_name"])

# --------------------------------------------------------------------------
# Login audit, leaderboard and awards
# --------------------------------------------------------------------------
def audit_scope(u):
    """SQL filter + params restricting what the user may see."""
    if u["role"] == "admin":
        return "1=1", []
    if u["role"] == "manager":
        return "u.branch_id=?", [u["branch_id"]]
    return "u.id=?", [u["id"]]


def audit_rows(u, f, t, user_id=None, mode="sessions"):
    scope, params = audit_scope(u)
    if user_id and u["role"] != "cashier":
        scope += " AND u.id=?"
        params.append(user_id)
    if mode == "daily":
        return q("SELECT date(ls.login_at) AS date, u.full_name AS staff, u.role, b.name AS branch, COUNT(ls.id) AS logins, "
                 "MIN(time(ls.login_at)) AS first_login, MAX(time(COALESCE(ls.logout_at, ls.last_seen_at, ls.login_at))) AS last_logout, "
                 "ROUND(SUM((julianday(COALESCE(ls.logout_at, ls.last_seen_at, ls.login_at)) - julianday(ls.login_at)) * 24), 2) AS hours, "
                 "(SELECT COUNT(*) FROM sales s WHERE s.created_by=u.id AND s.date=date(ls.login_at)) AS bills, "
                 "(SELECT COALESCE(SUM(grand_total),0) FROM sales s WHERE s.created_by=u.id AND s.date=date(ls.login_at)) AS sales "
                 "FROM login_sessions ls JOIN users u ON u.id=ls.user_id LEFT JOIN branches b ON b.id=ls.branch_id "
                 f"WHERE {scope} AND date(ls.login_at) BETWEEN ? AND ? GROUP BY date(ls.login_at), u.id ORDER BY date DESC, staff",
                 params + [f, t])
    return q("SELECT ls.*, u.full_name AS staff, u.role, b.name AS branch, c.code AS counter, "
             "(SELECT COUNT(*) FROM sales s WHERE s.login_session_id=ls.id) AS bills, "
             "(SELECT COALESCE(SUM(grand_total),0) FROM sales s WHERE s.login_session_id=ls.id) AS sales, "
             "ROUND((julianday(COALESCE(ls.logout_at, ls.last_seen_at, ls.login_at)) - julianday(ls.login_at)) * 24, 2) AS hours "
             "FROM login_sessions ls JOIN users u ON u.id=ls.user_id LEFT JOIN branches b ON b.id=ls.branch_id "
             "LEFT JOIN counters c ON c.id=ls.counter_id "
             f"WHERE {scope} AND date(ls.login_at) BETWEEN ? AND ? ORDER BY ls.id DESC", params + [f, t])


@app.route("/audit")
@login_required
def audit():
    u = current_user()
    f = request.args.get("from") or date.today().replace(day=1).isoformat()
    t = request.args.get("to") or date.today().isoformat()
    mode = request.args.get("mode", "sessions")
    user_id = request.args.get("user_id", type=int)
    rows = audit_rows(u, f, t, user_id, mode)
    if u["role"] == "admin":
        staff = q("SELECT id, full_name FROM users WHERE active=1 ORDER BY full_name")
    elif u["role"] == "manager":
        staff = q("SELECT id, full_name FROM users WHERE active=1 AND branch_id=? ORDER BY full_name", (u["branch_id"],))
    else:
        staff = []
    return render_template("audit.html", rows=rows, date_from=f, date_to=t, mode=mode, staff=staff, user_id=user_id)


@app.route("/audit/export")
@login_required
def audit_export():
    u = current_user()
    f, t = request.args.get("from"), request.args.get("to")
    mode = request.args.get("mode", "sessions")
    rows = audit_rows(u, f, t, request.args.get("user_id", type=int), mode)
    if mode == "daily":
        cols = [("date", "Date"), ("staff", "Staff"), ("role", "Role"), ("branch", "Branch"), ("logins", "Logins"), ("first_login", "First login"),
                ("last_logout", "Last logout"), ("hours", "Hours"), ("bills", "Bills"), ("sales", "Sales")]
    else:
        cols = [("staff", "Staff"), ("role", "Role"), ("branch", "Branch"), ("counter", "Counter"), ("login_at", "Login"), ("logout_at", "Logout"),
                ("last_seen_at", "Last active"), ("end_reason", "How ended"), ("hours", "Hours"), ("bills", "Bills"), ("sales", "Sales"), ("ip", "IP")]
    return export_response(f"Login audit ({'daily summary' if mode == 'daily' else 'each login'})", f"{f} to {t}", rows,
                           request.args.get("fmt", "xlsx"), cols, filename=f"login_audit_{mode}")


def period_range(key):
    today = date.today()
    if key == "today":
        return today.isoformat(), today.isoformat()
    if key == "week":
        return (today - timedelta(days=today.weekday())).isoformat(), today.isoformat()
    if key == "year":
        return today.replace(month=1, day=1).isoformat(), today.isoformat()
    return today.replace(day=1).isoformat(), today.isoformat()


@app.route("/leaderboard")
@login_required
def leaderboard():
    u = current_user()
    period = request.args.get("period", "month")
    if period == "custom" and request.args.get("from") and request.args.get("to"):
        f, t = request.args["from"], request.args["to"]
    else:
        f, t = period_range(period)
    if u["role"] == "admin":
        bsel = request.args.get("branch", "all")
        bid = None if bsel == "all" else int(bsel)
    else:
        bid = u["branch_id"]
    where, params = ("s.date BETWEEN ? AND ?", [f, t])
    if bid:
        where += " AND s.branch_id=?"
        params.append(bid)
    rows = q("SELECT u.id, u.full_name, u.role, b.name AS branch, COUNT(s.id) AS bills, COALESCE(SUM(s.grand_total),0) AS amount, "
             "ROUND(COALESCE(AVG(s.grand_total),0),2) AS avg_bill, "
             "(SELECT COALESCE(SUM(si.quantity),0) FROM sale_items si JOIN sales s2 ON s2.id=si.sale_id WHERE s2.created_by=u.id "
             " AND s2.date BETWEEN ? AND ?" + (" AND s2.branch_id=?" if bid else "") + ") AS items, "
             "COALESCE(SUM(s.credit_amount),0) AS credit FROM sales s JOIN users u ON u.id=s.created_by "
             "LEFT JOIN branches b ON b.id=s.branch_id "
             f"WHERE {where} GROUP BY u.id, s.branch_id ORDER BY amount DESC", [f, t] + ([bid] if bid else []) + params)
    awards = q("SELECT a.*, u.full_name AS staff, g.full_name AS giver, b.name AS branch FROM awards a JOIN users u ON u.id=a.user_id "
               "JOIN users g ON g.id=a.given_by LEFT JOIN branches b ON b.id=a.branch_id " +
               ("WHERE a.branch_id=? " if bid else "") + "ORDER BY a.id DESC LIMIT 20", [bid] if bid else [])
    my_awards = q("SELECT * FROM awards WHERE user_id=? ORDER BY id DESC", (u["id"],))
    branches_list = q("SELECT * FROM branches WHERE active=1 ORDER BY name") if u["role"] == "admin" else []
    return render_template("leaderboard.html", rows=rows, period=period, date_from=f, date_to=t, bid=bid,
                           branches_list=branches_list, awards=awards, my_awards=my_awards)


@app.route("/awards", methods=["POST"])
@role_required("admin", "manager")
def award_give():
    u = current_user()
    f = request.form
    winner = q("SELECT * FROM users WHERE id=? AND active=1", (f.get("user_id"),), one=True)
    title = f.get("title", "").strip()
    if not winner or not title:
        flash("Choose a staff member and give the prize a title.", "error")
        return redirect(request.referrer or url_for("leaderboard"))
    if u["role"] == "manager" and winner["branch_id"] != u["branch_id"]:
        flash("You can only award staff of your own branch.", "error")
        return redirect(url_for("leaderboard"))
    bid = winner["branch_id"] or session.get("branch_id")
    run("INSERT INTO awards(user_id,branch_id,given_by,title,prize,note,period_from,period_to) VALUES (?,?,?,?,?,?,?,?)",
        (winner["id"], bid, u["id"], title, f.get("prize", "").strip(), f.get("note", "").strip(), f.get("period_from"), f.get("period_to")))
    notify([winner["id"]], f"Congratulations! {title}",
           f"You have been awarded \"{title}\" by {u['full_name']}" + (f" for the period {f.get('period_from')} to {f.get('period_to')}" if f.get("period_from") else "")
           + (f". Prize: {f.get('prize').strip()}" if f.get("prize", "").strip() else "") + (f". {f.get('note').strip()}" if f.get("note", "").strip() else ""), u["id"])
    get_db().commit()
    flash(f"Award recorded for {winner['full_name']}.", "success")
    return redirect(url_for("leaderboard", period=request.args.get("period", "month")))


@app.route("/awards/<int:aid>/delete", methods=["POST"])
@role_required("admin")
def award_delete(aid):
    run("DELETE FROM awards WHERE id=?", (aid,))
    get_db().commit()
    flash("Award removed.", "success")
    return redirect(url_for("leaderboard"))

# --------------------------------------------------------------------------
# Activity log: records every action automatically (admin only can view)
# --------------------------------------------------------------------------
ACTION_LABELS = {
    "login": "Login", "logout": "Logout", "select_counter": "Selected counter", "switch_branch": "Switched branch",
    "change_password": "Changed own password",
    "api_sale_create": "Created bill", "sale_delete": "Cancelled bill", "sale_detail": "Printed bill",
    "customer_payment": "Received udhaar payment", "receipt_delete": "Cancelled receipt", "receipt_detail": "Printed receipt",
    "product_new": "Added product", "product_edit": "Edited product (prices/details)", "product_delete": "Removed product",
    "categories": "Added category", "category_delete": "Deleted category",
    "stock_adjust": "Adjusted stock", "stock_transfer": "Transferred stock", "purchase_new": "Recorded purchase",
    "suppliers": "Saved supplier", "customers": "Saved customer", "api_customer_create": "Added customer",
    "settings_page": "Changed shop settings", "branches": "Saved branch", "counters_save": "Saved counter", "users": "Saved user",
    "leave_policy": "Changed leave policy",
    "notice_new": "Sent notice", "notice_delete": "Deleted notice",
    "message_new": "Sent branch message", "message_thread": "Replied to branch message",
    "message_send_stock": "Sent stock for request", "message_status": "Changed request status",
    "leave_new": "Applied for leave", "leave_decide": "Decided leave", "leave_cancel": "Withdrew leave",
    "leave_document": "Leave document", "award_give": "Gave award", "award_delete": "Removed award",
    "user_reset_device": "Reset device lock", "offer_form": "Saved offer",
    "purchase_bill": "Supplier bill", "signature_generate": "Created digital signature", "branch_signature_save": "Saved branch signature", "hr_staff": "Saved staff profile", "hr_offer_letter": "Generated offer letter", "hr_id_card": "Generated ID card",
    "hr_salary_slip": "Generated salary slip", "document_view": "Viewed staff document", "document_delete": "Deleted staff document",
    "document_signed": "Signed offer letter copy", "offer_toggle": "Switched offer on/off", "offer_delete": "Deleted offer",
    "report_export": "Exported report", "audit_export": "Exported login audit", "activity_export": "Exported activity log",
    "sales_export": "Exported bills register", "receipts_export": "Exported receipts", "stock_export": "Exported stock list",
    "customers_export": "Exported customer list",
}
SKIP_ENDPOINTS = {"static", "api_product_search", "api_customer_search", "api_notices_unread", "api_branch_stock", "notice_mark_read", "signature_preview"}
LOGGED_GETS = {"logout", "switch_branch", "report_export", "audit_export", "activity_export", "sales_export", "receipts_export",
               "stock_export", "customers_export", "document_view", "purchase_bill", "document_signed"}
SECRET_KEYS = ("password", "old", "new", "confirm")


def _clean(d):
    out = {}
    for k, v in d.items():
        if any(x in k.lower() for x in SECRET_KEYS):
            continue
        if isinstance(v, (list, tuple)):
            v = list(v)
        elif isinstance(v, dict):
            v = _clean(v)
        out[k] = v
    return out


@app.after_request
def record_activity(response):
    try:
        ep = request.endpoint
        if not ep or ep in SKIP_ENDPOINTS:
            return response
        is_print = ep in ("sale_detail", "receipt_detail") and request.args.get("print") == "1"
        is_doc_view = ep == "leave_document" and request.method == "GET"
        if request.method == "GET" and not (ep in LOGGED_GETS or is_print or is_doc_view):
            return response
        if ep == "login" and request.method == "GET":
            return response
        details = {}
        if request.is_json:
            body = request.get_json(silent=True) or {}
            if isinstance(body, dict):
                body = dict(body)
                if "items" in body and isinstance(body["items"], list):
                    body["items"] = [{k: v for k, v in i.items() if k in ("product_id", "quantity", "unit_price", "unit_cost")} for i in body["items"]]
                details = _clean(body)
        elif request.form:
            details = _clean({k: (request.form.getlist(k) if len(request.form.getlist(k)) > 1 else request.form.get(k)) for k in request.form})
        if request.files:
            details["files"] = [f.filename for f in request.files.values() if f.filename]
        if request.args and request.method == "GET":
            details["args"] = dict(request.args)
        action = ACTION_LABELS.get(ep, ep.replace("_", " ").title())
        if getattr(g, "audit_note", None):
            details["note"] = g.audit_note
        if ep == "login":
            action = "Login" if response.status_code == 302 else ("Login BLOCKED (device)" if getattr(g, "audit_note", "").startswith("Device lock") else "Login FAILED")
        elif ep == "leave_document":
            action = "Uploaded leave document" if request.method == "POST" else "Viewed leave document"
        elif ep == "purchase_bill":
            action = "Attached supplier bill" if request.method == "POST" else "Viewed supplier bill"
        elif ep == "document_signed":
            action = "Uploaded signed offer letter" if request.method == "POST" else "Viewed signed offer letter"
        elif ep == "users" and not request.form.get("id"):
            action = "Created user"
        elif ep == "leave_decide":
            action = f"Leave {request.form.get('status', '')}"
        elif ep in ("suppliers", "customers", "branches", "counters_save") and not request.form.get("id"):
            action = action.replace("Saved", "Added")
        entity_id = None
        if request.view_args:
            entity_id = next((str(v) for v in request.view_args.values()), None)
        if response.is_json and response.status_code < 500 and response.content_length and response.content_length < 2000:
            try:
                j = response.get_json()
                if isinstance(j, dict):
                    entity_id = j.get("invoice_no") or (str(j.get("id")) if j.get("id") else entity_id)
                    if j.get("grand_total") is not None:
                        details["grand_total"] = j["grand_total"]
                    if j.get("error"):
                        details["error"] = j["error"]
            except Exception:
                pass
        u = getattr(g, "audit_user", None) or (current_user() if session.get("user_id") else None)
        bid_log = getattr(g, "audit_branch", None) or session.get("branch_id")
        username = u["username"] if u else (request.form.get("username") if ep == "login" else None)
        get_db().execute("INSERT INTO audit_log(user_id,username,role,branch_id,ip,method,endpoint,action,entity_id,details,status) "
                         "VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                         (u["id"] if u else None, username, u["role"] if u else None, bid_log, request.remote_addr,
                          request.method, ep, action, entity_id, json.dumps(details, default=str)[:2000], response.status_code))
        get_db().commit()
    except Exception as e:                     # logging must never break the app
        app.logger.warning("activity log failed: %s", e)
    return response


def activity_rows(f, t, user_id=None, branch_id_=None, text=None, limit=2000):
    where, params = ["date(a.at) BETWEEN ? AND ?"], [f, t]
    if user_id:
        where.append("a.user_id=?"); params.append(user_id)
    if branch_id_:
        where.append("a.branch_id=?"); params.append(branch_id_)
    if text:
        where.append("(a.action LIKE ? OR a.details LIKE ? OR a.entity_id LIKE ? OR a.username LIKE ?)")
        params += [f"%{text}%"] * 4
    return q("SELECT a.*, u.full_name, b.name AS branch FROM audit_log a LEFT JOIN users u ON u.id=a.user_id "
             "LEFT JOIN branches b ON b.id=a.branch_id WHERE " + " AND ".join(where) + f" ORDER BY a.id DESC LIMIT {int(limit)}", params)


@app.route("/activity")
@role_required("admin")
def activity():
    f = request.args.get("from") or date.today().isoformat()
    t = request.args.get("to") or date.today().isoformat()
    user_id = request.args.get("user_id", type=int)
    bid = request.args.get("branch_id", type=int)
    text = request.args.get("q", "").strip()
    rows = activity_rows(f, t, user_id, bid, text)
    for r in rows:
        pass
    return render_template("activity.html", rows=rows, date_from=f, date_to=t, user_id=user_id, bid=bid, text=text,
                           staff=q("SELECT id, full_name FROM users ORDER BY full_name"),
                           branches_list=q("SELECT * FROM branches ORDER BY name"))


@app.route("/activity/export")
@role_required("admin")
def activity_export():
    f = request.args.get("from") or date.today().isoformat()
    t = request.args.get("to") or date.today().isoformat()
    rows = activity_rows(f, t, request.args.get("user_id", type=int), request.args.get("branch_id", type=int),
                         request.args.get("q", "").strip(), limit=50000)
    rows = [dict(r, user=(r["full_name"] or r["username"]), result=("OK" if (r["status"] or 0) < 400 else f"Failed ({r['status']})")) for r in rows]
    cols = [("at", "Date/time"), ("user", "User"), ("role", "Role"), ("branch", "Branch"), ("action", "Action"), ("entity_id", "Reference"),
            ("details", "Details"), ("result", "Result"), ("ip", "IP")]
    return export_response("Activity log", f"{f} to {t}", rows, request.args.get("fmt", "xlsx"), cols, filename="activity_log")


@app.route("/sales/export")
@login_required
def sales_export():
    f = request.args.get("from") or date.today().isoformat()
    t = request.args.get("to") or date.today().isoformat()
    rows = q("SELECT s.invoice_no, s.date, s.created_at, co.code AS counter, COALESCE(c.name,'Walk-in') AS customer, u.full_name AS cashier, "
             "s.payment_mode, s.subtotal, s.discount, s.taxable, s.cgst, s.sgst, s.round_off, s.grand_total, s.paid_amount, s.credit_amount "
             "FROM sales s LEFT JOIN customers c ON c.id=s.customer_id LEFT JOIN users u ON u.id=s.created_by "
             "LEFT JOIN counters co ON co.id=s.counter_id WHERE s.branch_id=? AND s.date BETWEEN ? AND ? ORDER BY s.id", (branch_id(), f, t))
    return export_response("Bills register", f"{f} to {t}", rows, request.args.get("fmt", "xlsx"), filename="bills")


@app.route("/receipts/export")
@login_required
def receipts_export():
    f = request.args.get("from") or date.today().replace(day=1).isoformat()
    t = request.args.get("to") or date.today().isoformat()
    rows = q("SELECT p.receipt_no, p.date, c.name AS customer, c.phone, p.mode, p.note, p.amount, p.prev_balance, p.new_balance, u.full_name AS received_by "
             "FROM payments p JOIN customers c ON c.id=p.customer_id LEFT JOIN users u ON u.id=p.created_by "
             "WHERE p.branch_id=? AND p.date BETWEEN ? AND ? ORDER BY p.id", (branch_id(), f, t))
    return export_response("Udhaar payment receipts", f"{f} to {t}", rows, request.args.get("fmt", "xlsx"), filename="receipts")


@app.route("/stock/export")
@login_required
def stock_export():
    only_low = request.args.get("low") == "1"
    rows = q("SELECT p.name AS item, c.name AS category, COALESCE(s.quantity,0) AS quantity, p.unit, p.low_stock_level AS alert_level, "
             "p.purchase_price AS cost_per_unit, p.selling_price AS sell_price, p.gst_rate, p.hsn, "
             "COALESCE(s.quantity,0)*p.purchase_price AS stock_value FROM products p LEFT JOIN categories c ON c.id=p.category_id "
             "LEFT JOIN stock s ON s.product_id=p.id AND s.branch_id=? WHERE p.active=1 " +
             ("AND COALESCE(s.quantity,0) <= p.low_stock_level " if only_low else "") + "ORDER BY p.name", (branch_id(),))
    return export_response("Low stock list" if only_low else "Stock list", f"As on {date.today().isoformat()}", rows,
                           request.args.get("fmt", "xlsx"), filename="stock")


@app.route("/customers/export")
@login_required
def customers_export():
    rows = q("SELECT name, phone, address, credit_balance AS udhaar_due FROM customers ORDER BY credit_balance DESC, name")
    return export_response("Customers and udhaar due", f"As on {date.today().isoformat()}", rows, request.args.get("fmt", "xlsx"), filename="customers")

# --------------------------------------------------------------------------
# Analytics dashboard (admin, or users granted access)
# --------------------------------------------------------------------------
def analytics_allowed(u):
    return u and (u["role"] == "admin" or u["can_view_analytics"])


def _period(key, f_arg, t_arg):
    today = date.today()
    if key == "custom" and f_arg and t_arg:
        f, t = date.fromisoformat(f_arg), date.fromisoformat(t_arg)
    elif key == "today":
        f = t = today
    elif key == "7d":
        f, t = today - timedelta(days=6), today
    elif key == "30d":
        f, t = today - timedelta(days=29), today
    elif key == "year":
        f, t = today.replace(month=1, day=1), today
    elif key == "lastmonth":
        first = today.replace(day=1)
        t = first - timedelta(days=1)
        f = t.replace(day=1)
    else:
        key = "month"
        f, t = today.replace(day=1), today
    days = (t - f).days + 1
    pt = f - timedelta(days=1)
    pf = pt - timedelta(days=days - 1)
    return key, f.isoformat(), t.isoformat(), pf.isoformat(), pt.isoformat()


def analytics_data(f, t, pf, pt, bid):
    bw = " AND s.branch_id=? " if bid else " "
    bp = [bid] if bid else []

    def kpi(a, b):
        r = q("SELECT COUNT(*) AS bills, COALESCE(SUM(s.grand_total),0) AS sales, COALESCE(AVG(s.grand_total),0) AS avg_bill, "
              "COALESCE(SUM(s.credit_amount),0) AS credit, COALESCE(SUM(s.discount),0) AS discount, "
              "COUNT(DISTINCT s.customer_id) AS customers FROM sales s WHERE s.date BETWEEN ? AND ?" + bw, [a, b] + bp, one=True)
        r2 = q("SELECT COALESCE(SUM(si.quantity),0) AS items, COALESCE(SUM(si.total - si.cost_price*si.quantity),0) AS profit "
               "FROM sale_items si JOIN sales s ON s.id=si.sale_id WHERE s.date BETWEEN ? AND ?" + bw, [a, b] + bp, one=True)
        return dict(r) | dict(r2)

    cur, prev = kpi(f, t), kpi(pf, pt)

    def pct(a, b):
        return None if not b else round((a - b) / b * 100, 1)

    growth = {k: pct(cur[k], prev[k]) for k in cur}
    trend = q("SELECT s.date, b.name AS branch, SUM(s.grand_total) AS sales, COUNT(*) AS bills FROM sales s JOIN branches b ON b.id=s.branch_id "
              "WHERE s.date BETWEEN ? AND ?" + bw + "GROUP BY s.date, s.branch_id ORDER BY s.date", [f, t] + bp)
    prev_trend = q("SELECT s.date, SUM(s.grand_total) AS sales FROM sales s WHERE s.date BETWEEN ? AND ?" + bw + "GROUP BY s.date ORDER BY s.date",
                   [pf, pt] + bp)
    branches_cmp = q("SELECT b.name, COALESCE(SUM(s.grand_total),0) AS sales, COUNT(s.id) AS bills, "
                     "(SELECT COALESCE(SUM(s2.grand_total),0) FROM sales s2 WHERE s2.branch_id=b.id AND s2.date BETWEEN ? AND ?) AS prev_sales "
                     "FROM branches b LEFT JOIN sales s ON s.branch_id=b.id AND s.date BETWEEN ? AND ? WHERE b.active=1 GROUP BY b.id ORDER BY sales DESC",
                     [pf, pt, f, t])
    payments = q("SELECT s.payment_mode AS mode, SUM(s.grand_total) AS amount, COUNT(*) AS bills FROM sales s WHERE s.date BETWEEN ? AND ?" + bw +
                 "GROUP BY s.payment_mode ORDER BY amount DESC", [f, t] + bp)
    top_products = q("SELECT si.name, SUM(si.quantity) AS qty, SUM(si.total) AS sales, SUM(si.total - si.cost_price*si.quantity) AS profit "
                     "FROM sale_items si JOIN sales s ON s.id=si.sale_id WHERE s.date BETWEEN ? AND ?" + bw +
                     "GROUP BY si.product_id ORDER BY sales DESC LIMIT 10", [f, t] + bp)
    categories = q("SELECT COALESCE(c.name,'Uncategorised') AS category, SUM(si.total) AS sales FROM sale_items si JOIN sales s ON s.id=si.sale_id "
                   "JOIN products p ON p.id=si.product_id LEFT JOIN categories c ON c.id=p.category_id WHERE s.date BETWEEN ? AND ?" + bw +
                   "GROUP BY c.id ORDER BY sales DESC LIMIT 8", [f, t] + bp)
    hourly = q("SELECT CAST(strftime('%H', s.created_at) AS INTEGER) AS hour, COUNT(*) AS bills, SUM(s.grand_total) AS sales FROM sales s "
               "WHERE s.date BETWEEN ? AND ?" + bw + "GROUP BY hour ORDER BY hour", [f, t] + bp)
    weekday = q("SELECT CAST(strftime('%w', s.date) AS INTEGER) AS dow, COUNT(*) AS bills, SUM(s.grand_total) AS sales FROM sales s "
                "WHERE s.date BETWEEN ? AND ?" + bw + "GROUP BY dow ORDER BY dow", [f, t] + bp)
    first_month = (date.today().replace(day=1) - timedelta(days=330)).replace(day=1).isoformat()
    monthly = q("SELECT substr(s.date,1,7) AS month, SUM(s.grand_total) AS sales, COUNT(*) AS bills, "
                "(SELECT COALESCE(SUM(si.total - si.cost_price*si.quantity),0) FROM sale_items si JOIN sales s2 ON s2.id=si.sale_id "
                " WHERE substr(s2.date,1,7)=substr(s.date,1,7)" + (" AND s2.branch_id=?" if bid else "") + ") AS profit "
                "FROM sales s WHERE s.date >= ?" + bw + "GROUP BY month ORDER BY month", bp + [first_month] + bp)
    slow = q("SELECT p.name, SUM(st.quantity) AS qty, p.unit, SUM(st.quantity*p.purchase_price) AS value FROM products p JOIN stock st ON st.product_id=p.id "
             "WHERE p.active=1 AND st.quantity>0" + (" AND st.branch_id=?" if bid else "") +
             " AND p.id NOT IN (SELECT si.product_id FROM sale_items si JOIN sales s ON s.id=si.sale_id WHERE s.date BETWEEN ? AND ?" + bw + ") "
             "GROUP BY p.id ORDER BY value DESC LIMIT 10", bp + [f, t] + bp)
    top_customers = q("SELECT c.name, COUNT(*) AS bills, SUM(s.grand_total) AS sales, c.credit_balance FROM sales s JOIN customers c ON c.id=s.customer_id "
                      "WHERE s.date BETWEEN ? AND ?" + bw + "GROUP BY c.id ORDER BY sales DESC LIMIT 10", [f, t] + bp)
    udhaar = q("SELECT COALESCE(SUM(credit_balance),0) AS total, SUM(CASE WHEN credit_balance>0 THEN 1 ELSE 0 END) AS customers FROM customers", one=True)
    low_stock = q("SELECT COUNT(*) AS n FROM products p JOIN stock st ON st.product_id=p.id WHERE p.active=1 AND st.quantity<=p.low_stock_level" +
                  (" AND st.branch_id=?" if bid else ""), bp, one=True)["n"]
    stock_value = q("SELECT COALESCE(SUM(st.quantity*p.purchase_price),0) AS v FROM stock st JOIN products p ON p.id=st.product_id WHERE p.active=1" +
                    (" AND st.branch_id=?" if bid else ""), bp, one=True)["v"]
    cashiers = q("SELECT u.full_name, COUNT(*) AS bills, SUM(s.grand_total) AS sales FROM sales s JOIN users u ON u.id=s.created_by "
                 "WHERE s.date BETWEEN ? AND ?" + bw + "GROUP BY u.id ORDER BY sales DESC LIMIT 8", [f, t] + bp)

    # plain-language insights
    ins = []
    if prev["sales"]:
        g = growth["sales"]
        ins.append(f"Sales are {'up' if g >= 0 else 'down'} {abs(g)}% compared with the previous {((date.fromisoformat(t) - date.fromisoformat(f)).days + 1)} days.")
    if growth["avg_bill"] is not None:
        ins.append(f"Average bill is {inr(cur['avg_bill'])} ({'up' if growth['avg_bill'] >= 0 else 'down'} {abs(growth['avg_bill'])}%).")
    if trend:
        best = max(trend, key=lambda r: r["sales"])
        ins.append(f"Best day was {best['date']} ({inr(best['sales'])}{' at ' + best['branch'] if not bid else ''}).")
    if hourly:
        ph = max(hourly, key=lambda r: r["bills"])
        ins.append(f"Peak billing hour is {ph['hour']:02d}:00-{ph['hour'] + 1:02d}:00 with {ph['bills']} bills. Keep enough counters open then.")
    if weekday:
        names = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
        pw = max(weekday, key=lambda r: r["sales"])
        ins.append(f"{names[pw['dow']]} is the strongest day of the week.")
    if top_products:
        ins.append(f"Top seller: {top_products[0]['name']} ({inr(top_products[0]['sales'])}). Never let it run out.")
    if slow:
        ins.append(f"{len(slow)}+ items had no sale in this period but hold stock worth {inr(sum(r['value'] for r in slow))}. Consider offers or smaller orders.")
    if udhaar["total"]:
        ins.append(f"Udhaar outstanding is {inr(udhaar['total'])} from {udhaar['customers']} customers. Collect regularly to protect cash flow.")
    if cur["sales"] and payments:
        credit_share = next((r["amount"] for r in payments if r["mode"] == "Credit"), 0) / cur["sales"] * 100
        if credit_share > 20:
            ins.append(f"{credit_share:.0f}% of sales went on credit. That is high; set limits per customer.")
    if low_stock:
        ins.append(f"{low_stock} items are at or below their alert level.")

    def L(rows):
        return [dict(r) for r in rows]
    return {"kpi": cur, "prev": prev, "growth": growth, "trend": L(trend), "prev_trend": L(prev_trend), "branches": L(branches_cmp),
            "payments": L(payments), "top_products": L(top_products), "categories": L(categories), "hourly": L(hourly), "weekday": L(weekday),
            "monthly": L(monthly), "slow": L(slow), "top_customers": L(top_customers), "udhaar": dict(udhaar), "low_stock": low_stock,
            "stock_value": stock_value, "cashiers": L(cashiers), "insights": ins}


@app.route("/analytics")
@login_required
def analytics():
    u = current_user()
    if not analytics_allowed(u):
        flash("You do not have access to the analytics dashboard.", "error")
        return redirect(url_for("dashboard"))
    key, f, t, pf, pt = _period(request.args.get("period", "month"), request.args.get("from"), request.args.get("to"))
    bsel = request.args.get("branch", "all")
    bid = None if bsel == "all" else int(bsel)
    if u["role"] != "admin" and u["branch_id"]:
        bid = u["branch_id"]                     # non-admin viewers see their own branch
    data = analytics_data(f, t, pf, pt, bid)
    return render_template("analytics.html", d=data, period=key, date_from=f, date_to=t, prev_from=pf, prev_to=pt, bid=bid,
                           branches_list=q("SELECT * FROM branches WHERE active=1 ORDER BY name"), data_json=json.dumps(data, default=str))

# --------------------------------------------------------------------------
# HR: staff profiles, offer letters, ID cards, salary slips
# --------------------------------------------------------------------------
DOC_DIR = os.path.join(BASE_DIR, "uploads", "docs")
PHOTO_DIR = os.path.join(BASE_DIR, "uploads", "photos")
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def next_emp_code():
    prefix = get_setting("emp_id_prefix", "RKT").strip() or "RKT"
    nums = [int(r["emp_code"][len(prefix) + 1:]) for r in q("SELECT emp_code FROM staff_profiles WHERE emp_code LIKE ?", (prefix + "-%",))
            if r["emp_code"][len(prefix) + 1:].isdigit()]
    return f"{prefix}-{(max(nums) + 1 if nums else 1):03d}"


def designation_list():
    return [x.strip() for x in get_setting("designations", "").split(",") if x.strip()]


def signature_path():
    f = get_setting("signature_file", "")
    path = os.path.join(BASE_DIR, "uploads", f) if f else None
    return path if path and os.path.exists(path) else None


def signature_for(branch_id_):
    """(image path or None, signatory name, source) for a branch, falling back to the company-wide signature."""
    if branch_id_:
        b = q("SELECT signature_file, signatory_name FROM branches WHERE id=?", (branch_id_,), one=True)
        if b and b["signature_file"]:
            path = os.path.join(BASE_DIR, "uploads", b["signature_file"])
            if os.path.exists(path):
                return path, b["signatory_name"] or "", "branch"
    return signature_path(), get_setting("signatory_name", ""), "company"


def hr_can_manage(u, target):
    if u["role"] == "admin":
        return True
    return bool(u["role"] == "manager" and u["can_hr_docs"] and target["role"] == "cashier" and target["branch_id"] == u["branch_id"])


def hr_user(uid):
    return q("SELECT u.*, b.name AS branch_name, b.address AS branch_address, b.phone AS branch_phone FROM users u "
             "LEFT JOIN branches b ON b.id=u.branch_id WHERE u.id=?", (uid,), one=True)


def profile_of(uid):
    r = q("SELECT * FROM staff_profiles WHERE user_id=?", (uid,), one=True)
    return dict(r) if r else {"user_id": uid, "monthly_salary": 0, "hra": 0, "allowances": 0}


def _pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        if "Arial" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Arial", r"C:\Windows\Fonts\arial.ttf"))
            pdfmetrics.registerFont(TTFont("Arial-Bold", r"C:\Windows\Fonts\arialbd.ttf"))
        return "Arial", "Arial-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def _save_doc(uid, kind, title, period, pdf_bytes, data):
    os.makedirs(DOC_DIR, exist_ok=True)
    fname = f"{kind}_{uid}_{int(datetime.now().timestamp())}.pdf"
    with open(os.path.join(DOC_DIR, fname), "wb") as fh:
        fh.write(pdf_bytes)
    did = run("INSERT INTO staff_documents(user_id,kind,title,period,file,data,created_by) VALUES (?,?,?,?,?,?,?)",
              (uid, kind, title, period, fname, json.dumps(data, default=str), session["user_id"]))
    notify([uid], f"New document: {title}", "Open My documents in the menu to view or print it.", session["user_id"])
    get_db().commit()
    return did


def logo_drawing(size_pt=42):
    """The RK logo mark as a reportlab drawing (None if it cannot be loaded)."""
    try:
        from svglib.svglib import svg2rlg
        d = svg2rlg(os.path.join(BASE_DIR, "static", "logo-mark.svg"))
        k = size_pt / 100.0
        d.scale(k, k)
        d.width, d.height = size_pt, size_pt
        return d
    except Exception:
        return None


def _company_header(story, styles, font, bold):
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import ParagraphStyle
    st = all_settings()
    name = Paragraph(st.get("shop_name", ""), ParagraphStyle("c1", fontName=bold, fontSize=18, leading=22))
    sub = " | ".join(x for x in [st.get("shop_tagline"), st.get("shop_address"), ("Phone: " + st["shop_phone"]) if st.get("shop_phone") else "",
                                 ("GSTIN: " + st["shop_gstin"]) if st.get("shop_gstin") else ""] if x)
    subp = Paragraph(sub, ParagraphStyle("c2", fontName=font, fontSize=9, textColor="#444444", leading=12))
    logo = logo_drawing(42)
    if logo:
        t = Table([[logo, [name, subp]]], colWidths=[50, None])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0)]))
        story.append(t)
    else:
        story += [name, subp]
    story.append(Spacer(1, 10))


def build_offer_letter(user, prof):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    font, bold = _pdf_fonts()
    st = all_settings()
    salary = float(prof.get("monthly_salary") or 0) + float(prof.get("hra") or 0) + float(prof.get("allowances") or 0)
    fields = {"name": user["full_name"], "designation": prof.get("designation") or user["role"].title(), "branch": user["branch_name"] or "Head Office",
              "joining_date": user["joining_date"] or date.today().isoformat(), "salary": inr(salary, symbol=False),
              "salary_words": number_to_words(salary), "company": st.get("shop_name", ""), "date": date.today().strftime("%d %B %Y")}
    body = st.get("offer_letter_template", "")
    for k, v in fields.items():
        body = body.replace("{" + k + "}", str(v))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title="Offer Letter")
    normal = ParagraphStyle("n", fontName=font, fontSize=11, leading=16)
    story = []
    _company_header(story, None, font, bold)
    story.append(HRFlowable(width="100%", thickness=1))
    story.append(Spacer(1, 14))
    story.append(Paragraph(f"Date: {fields['date']}", normal))
    story.append(Spacer(1, 10))
    story.append(Paragraph("To,<br/><b>" + fields["name"] + "</b><br/>" + (prof.get("address") or "").replace("\n", "<br/>"), normal))
    story.append(Spacer(1, 12))
    story.append(Paragraph("<b>Subject: Offer of Employment</b>", normal))
    story.append(Spacer(1, 10))
    for para in body.split("\n\n"):
        story.append(Paragraph(para.replace("\n", "<br/>"), normal))
        story.append(Spacer(1, 8))
    story.append(Spacer(1, 24))
    story.append(Paragraph(f"For <b>{fields['company']}</b>", normal))
    sig, sname, _ = signature_for(user["branch_id"])
    if sig:
        from reportlab.platypus import Image as RLImage
        try:
            story.append(RLImage(sig, width=40 * mm, height=16 * mm, kind="proportional", hAlign="LEFT"))
        except Exception:
            story.append(Spacer(1, 30))
    else:
        story.append(Spacer(1, 30))
    story.append(Paragraph("______________________<br/>Authorised Signatory" + (("<br/>" + sname) if sname else ""), normal))
    story.append(Spacer(1, 30))
    story.append(Paragraph("I accept the above offer.<br/><br/><br/>______________________<br/>" + fields["name"] + " &nbsp;&nbsp; Date: ___________", normal))
    doc.build(story)
    return buf.getvalue(), fields


def build_id_card(user, prof):
    """Portrait ID card (54 x 85.6 mm), front and back side by side on A4 with cut lines."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from reportlab.lib import colors
    from reportlab.graphics import renderPDF
    font, bold = _pdf_fonts()
    st = all_settings()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle("ID Card")
    W, H = 54 * mm, 85.6 * mm
    teal, teal_dark, orange = colors.HexColor("#0f766e"), colors.HexColor("#0b5c56"), colors.HexColor("#f59e0b")
    years = int(st.get("id_card_validity_years") or 2)
    vt = prof.get("id_valid_till") or date.today().replace(year=date.today().year + years).isoformat()
    valid = date.fromisoformat(vt).strftime("%d/%m/%Y")
    emp = prof.get("emp_code") or f"EMP{user['id']:04d}"
    company = st.get("shop_name", "")
    sig, _sn, _ = signature_for(user["branch_id"])

    def rounded_card(x, y):
        c.setDash(3, 3); c.setStrokeColor(colors.grey); c.setLineWidth(0.5)
        c.roundRect(x, y, W, H, 3 * mm); c.setDash(); c.setLineWidth(1)

    def header(x, y, title):
        c.setFillColor(teal); c.roundRect(x, y + H - 17 * mm, W, 17 * mm, 3 * mm, fill=1, stroke=0)
        c.rect(x, y + H - 17 * mm, W, 6 * mm, fill=1, stroke=0)          # square off the bottom edge of the band
        logo = logo_drawing(11 * mm)
        if logo:
            renderPDF.draw(logo, c, x + 3 * mm, y + H - 14.5 * mm)
        c.setFillColor(colors.white)
        c.setFont(bold, 9.5 if len(company) <= 14 else 8); c.drawString(x + 16 * mm, y + H - 8.5 * mm, company[:22])
        c.setFont(font, 5.5); c.drawString(x + 16 * mm, y + H - 12 * mm, (title or st.get("shop_tagline") or "")[:34])
        c.setFillColor(orange); c.rect(x, y + H - 17.8 * mm, W, 0.8 * mm, fill=1, stroke=0)

    def fit(text, size, maxw):
        while size > 5 and c.stringWidth(text, bold, size) > maxw:
            size -= 0.5
        return size

    # ---------------- front ----------------
    x0, y0 = 30 * mm, A4[1] - 30 * mm - H
    c.setFillColor(colors.white); c.roundRect(x0, y0, W, H, 3 * mm, fill=1, stroke=0)
    header(x0, y0, None)
    # photo
    pw, ph = 24 * mm, 26.5 * mm
    px, py = x0 + (W - pw) / 2, y0 + H - 20 * mm - ph
    c.setFillColor(colors.HexColor("#e6f4f1")); c.roundRect(px - 1.2 * mm, py - 1.2 * mm, pw + 2.4 * mm, ph + 2.4 * mm, 2 * mm, fill=1, stroke=0)
    photo = os.path.join(PHOTO_DIR, prof["photo"]) if prof.get("photo") else None
    drawn = False
    if photo and os.path.exists(photo):
        try:
            c.drawImage(photo, px, py, pw, ph, preserveAspectRatio=True, anchor="c")
            drawn = True
        except Exception:
            drawn = False
    if not drawn:
        c.setFillColor(colors.HexColor("#cbd5e1")); c.rect(px, py, pw, ph, fill=1, stroke=0)
        c.setFillColor(colors.white); c.setFont(font, 7); c.drawCentredString(px + pw / 2, py + ph / 2 - 2, "PHOTO")
    # name + designation
    c.setFillColor(teal_dark)
    name = user["full_name"]
    c.setFont(bold, fit(name, 10.5, W - 6 * mm)); c.drawCentredString(x0 + W / 2, py - 5.5 * mm, name)
    c.setFillColor(orange); c.setFont(bold, 6.5)
    c.drawCentredString(x0 + W / 2, py - 9.5 * mm, (prof.get("designation") or user["role"].title()).upper()[:30])
    # details
    c.setFillColor(colors.black)
    rows = [("ID", emp), ("Branch", user["branch_name"] or "Head Office"), ("Phone", prof.get("phone") or "-"),
            ("Blood", prof.get("blood_group") or "-"), ("Joined", user["joining_date"] or "-")]
    yy = py - 14 * mm
    for n_, (k, v) in enumerate(rows):
        c.setFont(bold, 6); c.drawString(x0 + 5 * mm, yy, k)
        c.setFont(font, 6.3); c.drawString(x0 + 17 * mm, yy, str(v)[:26])
        if n_ < len(rows) - 1:
            c.setStrokeColor(colors.HexColor("#e5e7eb")); c.setLineWidth(0.3); c.line(x0 + 5 * mm, yy - 1.3 * mm, x0 + W - 5 * mm, yy - 1.3 * mm)
        yy -= 4.2 * mm
    # footer band
    c.setFillColor(teal); c.rect(x0, y0, W, 6 * mm, fill=1, stroke=0)
    c.setFillColor(colors.white); c.setFont(font, 5.5)
    c.drawString(x0 + 3 * mm, y0 + 2.3 * mm, f"Valid till {valid}")
    c.drawRightString(x0 + W - 3 * mm, y0 + 2.3 * mm, "Authorised Signatory")
    if sig:
        try:
            c.drawImage(sig, x0 + W - 20 * mm, y0 + 6.3 * mm, 16 * mm, 6 * mm, preserveAspectRatio=True, anchor="se", mask="auto")
        except Exception:
            pass
    rounded_card(x0, y0)

    # ---------------- back ----------------
    x1 = x0 + W + 20 * mm
    c.setFillColor(colors.white); c.roundRect(x1, y0, W, H, 3 * mm, fill=1, stroke=0)
    header(x1, y0, "EMPLOYEE IDENTITY CARD")
    c.setFillColor(colors.black)
    lines = [(bold, 6.3, "Terms of use"),
             (font, 5.8, f"This card is the property of {company}."), (font, 5.8, "Carry it during duty hours and show it on request."),
             (font, 5.8, "Return it when leaving service. Loss must be"), (font, 5.8, "reported to the manager immediately."),
             (None, 0, ""),
             (bold, 6.3, "Emergency contact"), (font, 5.8, (prof.get("emergency_contact") or "-")[:40]),
             (None, 0, ""),
             (bold, 6.3, "Holder's address"), (font, 5.8, (prof.get("address") or "-").replace("\n", ", ")[:44]),
             (None, 0, ""),
             (bold, 6.3, "If found, please return to"), (font, 5.8, (st.get("shop_address") or "")[:44]),
             (font, 5.8, f"Phone: {st.get('shop_phone') or '-'}")]
    yy = y0 + H - 23 * mm
    for fnt, size, txt in lines:
        if fnt:
            c.setFont(fnt, size); c.drawString(x1 + 4 * mm, yy, txt)
            yy -= 3.6 * mm
        else:
            yy -= 1.8 * mm
    c.setFont(font, 5.5); c.drawString(x1 + 4 * mm, y0 + 9 * mm, "Holder's signature")
    c.setStrokeColor(colors.black); c.setLineWidth(0.4); c.line(x1 + 4 * mm, y0 + 8 * mm, x1 + 30 * mm, y0 + 8 * mm)
    c.setFillColor(teal); c.rect(x1, y0, W, 6 * mm, fill=1, stroke=0)
    c.setFillColor(colors.white); c.setFont(font, 5.5); c.drawCentredString(x1 + W / 2, y0 + 2.3 * mm, f"ID {emp}  |  {company[:20]}")
    rounded_card(x1, y0)

    # labels and instructions
    c.setFillColor(colors.grey); c.setFont(font, 8)
    c.drawCentredString(x0 + W / 2, y0 - 6 * mm, "FRONT"); c.drawCentredString(x1 + W / 2, y0 - 6 * mm, "BACK")
    c.drawString(30 * mm, y0 - 16 * mm, "Print at 100% scale on card stock or laminate. Standard card size 54 x 85.6 mm (portrait).")
    c.drawString(30 * mm, y0 - 21 * mm, "Cut along the dashed lines and stick the two sides back to back.")
    c.showPage(); c.save()
    return buf.getvalue(), {"emp_code": emp, "valid_till": valid}


def month_days(ym):
    import calendar
    y, m = int(ym[:4]), int(ym[5:7])
    return calendar.monthrange(y, m)[1]


def leave_days_in_month(uid, ym, unpaid=True):
    """Approved leave days falling inside the month (Unpaid type only when unpaid=True, other types otherwise)."""
    y, m = int(ym[:4]), int(ym[5:7])
    first = date(y, m, 1)
    last = date(y, m, month_days(ym))
    rows = q("SELECT from_date, to_date, leave_type FROM leaves WHERE user_id=? AND status='approved' AND from_date<=? AND to_date>=?",
             (uid, last.isoformat(), first.isoformat()))
    total = 0
    for r in rows:
        if (r["leave_type"] == "Unpaid") != unpaid:
            continue
        a, b = max(date.fromisoformat(r["from_date"]), first), min(date.fromisoformat(r["to_date"]), last)
        total += (b - a).days + 1
    return total


def build_salary_slip(user, prof, ym, fig):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    font, bold = _pdf_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title="Salary Slip")
    n = ParagraphStyle("n", fontName=font, fontSize=9.5, leading=13)
    b = ParagraphStyle("b", parent=n, fontName=bold)
    story = []
    _company_header(story, None, font, bold)
    title = f"Salary Slip for {MONTHS[int(ym[5:7]) - 1]} {ym[:4]}"
    story.append(Paragraph(title, ParagraphStyle("t", fontName=bold, fontSize=14, leading=18, alignment=1)))
    story.append(Spacer(1, 8))
    info = [[Paragraph("Employee", b), Paragraph(user["full_name"], n), Paragraph("Employee ID", b), Paragraph(fig["emp_code"], n)],
            [Paragraph("Designation", b), Paragraph(prof.get("designation") or user["role"].title(), n), Paragraph("Branch", b), Paragraph(user["branch_name"] or "Head Office", n)],
            [Paragraph("Date of joining", b), Paragraph(user["joining_date"] or "-", n), Paragraph("PAN", b), Paragraph(prof.get("pan") or "-", n)],
            [Paragraph("Bank", b), Paragraph(prof.get("bank_name") or "-", n), Paragraph("Account no", b), Paragraph(prof.get("account_no") or "-", n)],
            [Paragraph("Days in month", b), Paragraph(str(fig["days_in_month"]), n), Paragraph("Paid leave / Unpaid leave", b), Paragraph(f"{fig['paid_leave']} / {fig['lop_days']}", n)],
            [Paragraph("Payable days", b), Paragraph(str(fig["payable_days"]), n), Paragraph("Pay date", b), Paragraph(fig["pay_date"], n)]]
    t = Table(info, colWidths=[32 * mm, 58 * mm, 40 * mm, 48 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9")),
                           ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#f1f5f9")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story.append(t)
    story.append(Spacer(1, 10))
    earn = [("Basic salary", fig["basic"]), ("House rent allowance", fig["hra"]), ("Other allowances", fig["allowances"]),
            ("Incentive / bonus", fig["incentive"]), ("Overtime", fig["overtime"])]
    ded = [("Provident fund", fig["pf"]), ("ESI", fig["esi"]), ("Professional tax", fig["ptax"]), ("Advance recovery", fig["advance"]),
           (f"Unpaid leave ({fig['lop_days']} day(s))", fig["lop_amount"]), ("Other deductions", fig["other_ded"])]
    rows = [[Paragraph("Earnings", b), Paragraph("Amount", b), Paragraph("Deductions", b), Paragraph("Amount", b)]]
    for i in range(max(len(earn), len(ded))):
        e = earn[i] if i < len(earn) else ("", None)
        d = ded[i] if i < len(ded) else ("", None)
        rows.append([Paragraph(e[0], n), Paragraph(inr(e[1], False) if e[1] is not None else "", n),
                     Paragraph(d[0], n), Paragraph(inr(d[1], False) if d[1] is not None else "", n)])
    rows.append([Paragraph("Gross earnings", b), Paragraph(inr(fig["gross"], False), b), Paragraph("Total deductions", b), Paragraph(inr(fig["deductions"], False), b)])
    t2 = Table(rows, colWidths=[55 * mm, 34 * mm, 55 * mm, 34 * mm])
    t2.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6f4f1")),
                            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f1f5f9")), ("ALIGN", (1, 1), (1, -1), "RIGHT"), ("ALIGN", (3, 1), (3, -1), "RIGHT")]))
    story.append(t2)
    story.append(Spacer(1, 10))
    t3 = Table([[Paragraph("NET PAY", ParagraphStyle("np", fontName=bold, fontSize=12)), Paragraph(inr(fig["net"]), ParagraphStyle("npv", fontName=bold, fontSize=13, alignment=2))],
                [Paragraph("Rupees " + number_to_words(fig["net"]) + " only", n), ""]], colWidths=[120 * mm, 58 * mm])
    t3.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1, colors.black), ("SPAN", (0, 1), (1, 1)), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#fffbeb"))]))
    story.append(t3)
    if fig.get("note"):
        story.append(Spacer(1, 6)); story.append(Paragraph("Note: " + fig["note"], n))
    story.append(Spacer(1, 14))
    sig, sname, _ = signature_for(user["branch_id"])
    sig_cell = ""
    if sig:
        from reportlab.platypus import Image as RLImage
        try:
            sig_cell = RLImage(sig, width=36 * mm, height=14 * mm, kind="proportional", hAlign="RIGHT")
        except Exception:
            sig_cell = ""
    story.append(Table([["", sig_cell], [Paragraph("Employee signature", n), Paragraph("Authorised signatory" + (f"<br/>{sname}" if sname else ""), ParagraphStyle("r", parent=n, alignment=2))]],
                       colWidths=[89 * mm, 89 * mm], rowHeights=[16 * mm, None]))
    story.append(Spacer(1, 6))
    story.append(Paragraph(f"Generated on {datetime.now().strftime('%d-%m-%Y %I:%M %p')}. This is a computer generated slip.",
                           ParagraphStyle("f", parent=n, fontSize=7.5, textColor="#666666")))
    doc.build(story)
    return buf.getvalue()


@app.route("/hr")
@login_required
def hr():
    u = current_user()
    if not (u["role"] == "admin" or (u["role"] == "manager" and u["can_hr_docs"])):
        flash("You do not have permission for that page.", "error")
        return redirect(url_for("dashboard"))
    base = ("SELECT u.id, u.username, u.full_name, u.role, u.joining_date, u.active, b.name AS branch_name, p.emp_code, p.designation, p.monthly_salary, "
            "p.photo, (SELECT COUNT(*) FROM staff_documents d WHERE d.user_id=u.id) AS docs, "
            "(SELECT COUNT(*) FROM staff_documents d WHERE d.user_id=u.id AND d.kind='offer_letter') AS offers, "
            "(SELECT COUNT(*) FROM staff_documents d WHERE d.user_id=u.id AND d.kind='offer_letter' AND d.signed_file IS NOT NULL) AS offers_signed, "
            "(SELECT MAX(period) FROM staff_documents d WHERE d.user_id=u.id AND d.kind='salary_slip') AS last_slip "
            "FROM users u LEFT JOIN branches b ON b.id=u.branch_id LEFT JOIN staff_profiles p ON p.user_id=u.id ")
    rows = q(base + "ORDER BY u.active DESC, b.name, u.full_name") if u["role"] == "admin" else \
        q(base + "WHERE u.role='cashier' AND u.branch_id=? ORDER BY u.active DESC, u.full_name", (u["branch_id"],))
    return render_template("hr.html", rows=rows)


@app.route("/hr/<int:uid>", methods=["GET", "POST"])
@login_required
def hr_staff(uid):
    u = current_user()
    target = hr_user(uid) or abort(404)
    if not hr_can_manage(u, target):
        abort(403)
    if request.method == "POST":
        f = request.form
        photo = None
        file = request.files.get("photo")
        if file and file.filename:
            ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
            if ext not in ("jpg", "jpeg", "png"):
                flash("Photo must be JPG or PNG.", "error")
                return redirect(url_for("hr_staff", uid=uid))
            os.makedirs(PHOTO_DIR, exist_ok=True)
            photo = f"u{uid}_{int(datetime.now().timestamp())}.{ext}"
            file.save(os.path.join(PHOTO_DIR, photo))
        current = profile_of(uid)
        vals = (current.get("emp_code") or next_emp_code(), f.get("designation", "").strip(), f.get("phone", "").strip(), f.get("address", "").strip(),
                f.get("dob") or None, f.get("blood_group", "").strip(), float(f.get("monthly_salary") or 0), float(f.get("hra") or 0),
                float(f.get("allowances") or 0), f.get("bank_name", "").strip(), f.get("account_no", "").strip(), f.get("ifsc", "").strip(),
                f.get("pan", "").strip().upper(), f.get("emergency_contact", "").strip())
        run("INSERT INTO staff_profiles(user_id,emp_code,designation,phone,address,dob,blood_group,monthly_salary,hra,allowances,bank_name,account_no,ifsc,pan,"
            "emergency_contact,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,datetime('now','localtime')) "
            "ON CONFLICT(user_id) DO UPDATE SET emp_code=excluded.emp_code, designation=excluded.designation, phone=excluded.phone, address=excluded.address, "
            "dob=excluded.dob, blood_group=excluded.blood_group, monthly_salary=excluded.monthly_salary, hra=excluded.hra, allowances=excluded.allowances, "
            "bank_name=excluded.bank_name, account_no=excluded.account_no, ifsc=excluded.ifsc, pan=excluded.pan, emergency_contact=excluded.emergency_contact, "
            "updated_at=excluded.updated_at", (uid,) + vals)
        if photo:
            run("UPDATE staff_profiles SET photo=? WHERE user_id=?", (photo, uid))
        if f.get("joining_date"):
            run("UPDATE users SET joining_date=? WHERE id=?", (f["joining_date"], uid))
        get_db().commit()
        flash("Profile saved.", "success")
        return redirect(url_for("hr_staff", uid=uid))
    docs = q("SELECT d.*, c.full_name AS creator FROM staff_documents d LEFT JOIN users c ON c.id=d.created_by WHERE d.user_id=? ORDER BY d.id DESC", (uid,))
    years = int(get_setting("id_card_validity_years", "2") or 2)
    default_valid = date.today().replace(year=date.today().year + years).isoformat()
    return render_template("hr_staff.html", t=target, p=profile_of(uid), docs=docs, this_month=date.today().strftime("%Y-%m"),
                           designations=designation_list(), default_valid=default_valid)


@app.route("/hr/<int:uid>/offer-letter", methods=["POST"])
@login_required
def hr_offer_letter(uid):
    u = current_user()
    target = hr_user(uid) or abort(404)
    if not hr_can_manage(u, target):
        abort(403)
    pdf, fields = build_offer_letter(target, profile_of(uid))
    did = _save_doc(uid, "offer_letter", f"Offer letter - {target['full_name']}", None, pdf, fields)
    flash("Offer letter generated.", "success")
    return redirect(url_for("document_view", did=did))


@app.route("/hr/<int:uid>/id-card", methods=["POST"])
@login_required
def hr_id_card(uid):
    u = current_user()
    target = hr_user(uid) or abort(404)
    if not hr_can_manage(u, target):
        abort(403)
    valid = request.form.get("valid_till")
    if valid:
        try:
            date.fromisoformat(valid)
        except ValueError:
            flash("Enter a valid date for the card validity.", "error")
            return redirect(url_for("hr_staff", uid=uid))
        run("INSERT INTO staff_profiles(user_id,emp_code,id_valid_till) VALUES (?,?,?) ON CONFLICT(user_id) DO UPDATE SET id_valid_till=excluded.id_valid_till",
            (uid, next_emp_code(), valid))
        get_db().commit()
    pdf, data = build_id_card(target, profile_of(uid))
    did = _save_doc(uid, "id_card", f"ID card - {target['full_name']}", None, pdf, data)
    flash("ID card generated.", "success")
    return redirect(url_for("document_view", did=did))


@app.route("/hr/<int:uid>/salary-slip", methods=["GET", "POST"])
@login_required
def hr_salary_slip(uid):
    u = current_user()
    target = hr_user(uid) or abort(404)
    if not hr_can_manage(u, target):
        abort(403)
    prof = profile_of(uid)
    ym = (request.form.get("period") or request.args.get("period") or date.today().strftime("%Y-%m"))[:7]
    try:
        dim = month_days(ym)
    except (ValueError, IndexError):
        flash("Choose a valid month.", "error")
        return redirect(url_for("hr_staff", uid=uid))
    lop_auto = leave_days_in_month(uid, ym, unpaid=True)
    paid_leave = leave_days_in_month(uid, ym, unpaid=False)
    if request.method == "POST" and request.form.get("generate"):
        f = request.form
        g = lambda k: round(float(f.get(k) or 0), 2)
        fig = {"emp_code": prof.get("emp_code") or f"EMP{uid:04d}", "days_in_month": dim, "lop_days": int(float(f.get("lop_days") or 0)),
               "paid_leave": paid_leave, "basic": g("basic"), "hra": g("hra"), "allowances": g("allowances"), "incentive": g("incentive"),
               "overtime": g("overtime"), "pf": g("pf"), "esi": g("esi"), "ptax": g("ptax"), "advance": g("advance"), "other_ded": g("other_ded"),
               "note": f.get("note", "").strip(), "pay_date": f.get("pay_date") or date.today().isoformat()}
        fixed = fig["basic"] + fig["hra"] + fig["allowances"]
        fig["lop_amount"] = round(fixed / dim * fig["lop_days"], 2)
        fig["payable_days"] = dim - fig["lop_days"]
        fig["gross"] = round(fixed + fig["incentive"] + fig["overtime"], 2)
        fig["deductions"] = round(fig["pf"] + fig["esi"] + fig["ptax"] + fig["advance"] + fig["other_ded"] + fig["lop_amount"], 2)
        fig["net"] = round(fig["gross"] - fig["deductions"], 2)
        pdf = build_salary_slip(target, prof, ym, fig)
        did = _save_doc(uid, "salary_slip", f"Salary slip - {MONTHS[int(ym[5:7]) - 1]} {ym[:4]} - {target['full_name']}", ym, pdf, fig)
        flash("Salary slip generated.", "success")
        return redirect(url_for("document_view", did=did))
    existing = q("SELECT id FROM staff_documents WHERE user_id=? AND kind='salary_slip' AND period=?", (uid, ym), one=True)
    return render_template("salary_form.html", t=target, p=prof, period=ym, days=dim, lop_auto=lop_auto, paid_leave=paid_leave,
                           existing=existing, months=MONTHS)


@app.route("/documents/<int:did>")
@login_required
def document_view(did):
    u = current_user()
    d = q("SELECT * FROM staff_documents WHERE id=?", (did,), one=True) or abort(404)
    target = hr_user(d["user_id"])
    if not (d["user_id"] == u["id"] or (target and hr_can_manage(u, target))):
        abort(403)
    return send_from_directory(DOC_DIR, d["file"], as_attachment=request.args.get("download") == "1",
                               download_name=secure_filename(d["title"]) + ".pdf", mimetype="application/pdf")


@app.route("/documents/<int:did>/signed", methods=["GET", "POST"])
@login_required
def document_signed(did):
    """Signed copy of an offer letter: the staff member (or HR) uploads it; HR and the staff member can view it."""
    u = current_user()
    d = q("SELECT * FROM staff_documents WHERE id=?", (did,), one=True) or abort(404)
    target = hr_user(d["user_id"])
    is_owner = d["user_id"] == u["id"]
    if not (is_owner or (target and hr_can_manage(u, target))):
        abort(403)
    if request.method == "POST":
        file = request.files.get("signed")
        if not file or not file.filename:
            flash("Choose the signed copy first.", "error")
            return redirect(request.referrer or url_for("my_documents"))
        ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
        if ext not in ALLOWED_DOCS:
            flash("Signed copy must be a PDF, JPG or PNG.", "error")
            return redirect(request.referrer or url_for("my_documents"))
        os.makedirs(DOC_DIR, exist_ok=True)
        stored = f"signed_{did}_{int(datetime.now().timestamp())}.{ext}"
        file.save(os.path.join(DOC_DIR, stored))
        if d["signed_file"]:
            try:
                os.remove(os.path.join(DOC_DIR, d["signed_file"]))
            except OSError:
                pass
        run("UPDATE staff_documents SET signed_file=?, signed_name=?, signed_at=datetime('now','localtime'), signed_by=? WHERE id=?",
            (stored, file.filename, u["id"], did))
        who = [r["id"] for r in q("SELECT id FROM users WHERE role='admin' AND active=1")] + [d["created_by"]]
        if is_owner:
            notify(who, f"Signed copy uploaded: {d['title']}", f"{u['full_name']} uploaded the signed copy. Open Staff HR to view it.", u["id"])
        else:
            notify([d["user_id"]], f"Signed copy filed: {d['title']}", "The office has filed your signed copy. You can view it under My documents.", u["id"])
        get_db().commit()
        flash("Signed copy saved.", "success")
        return redirect(request.referrer or url_for("my_documents"))
    if not d["signed_file"]:
        abort(404)
    return send_from_directory(DOC_DIR, d["signed_file"], as_attachment=False, download_name=d["signed_name"])


@app.route("/documents/<int:did>/delete", methods=["POST"])
@role_required("admin")
def document_delete(did):
    d = q("SELECT * FROM staff_documents WHERE id=?", (did,), one=True) or abort(404)
    try:
        os.remove(os.path.join(DOC_DIR, d["file"]))
    except OSError:
        pass
    run("DELETE FROM staff_documents WHERE id=?", (did,))
    get_db().commit()
    flash("Document deleted.", "success")
    return redirect(url_for("hr_staff", uid=d["user_id"]))


@app.route("/photos/<path:fname>")
@login_required
def photo_file(fname):
    return send_from_directory(PHOTO_DIR, fname)


@app.route("/my-documents")
@login_required
def my_documents():
    u = current_user()
    docs = q("SELECT d.*, c.full_name AS creator FROM staff_documents d LEFT JOIN users c ON c.id=d.created_by WHERE d.user_id=? ORDER BY d.id DESC", (u["id"],))
    return render_template("my_documents.html", docs=docs, p=profile_of(u["id"]), t=hr_user(u["id"]))


if __name__ == "__main__":
    if init_db():
        print("New database created. Login with  admin / admin123")
    print("RK Traders running at  http://localhost:5000   (press Ctrl+C to stop)")
    app.run(host="0.0.0.0", port=5000, debug=False)
